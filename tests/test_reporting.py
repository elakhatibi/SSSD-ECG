import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from care_ecg.reporting import export_results_workbook


class ReportingTests(unittest.TestCase):
    def test_excel_report_separates_text_mapping_from_direct_diagnosis(self):
        rows = [
            {
                "record_id": "r1",
                "expected": "NORM",
                "diagnosis": "NORM",
                "rationale": "Model-predicted diagnosis: NORM.",
                "probabilities": {"NORM": 0.9, "MI": 0.1},
            },
            {
                "record_id": "r2",
                "expected": "MI",
                "diagnosis": "MI",
                "rationale": "Normal ECG is considered before myocardial infarction.",
                "probabilities": {"NORM": 0.1, "MI": 0.9},
            },
        ]
        with tempfile.TemporaryDirectory() as directory:
            predictions = Path(directory) / "predictions.jsonl"
            output = Path(directory) / "results.xlsx"
            predictions.write_text(
                "\n".join(json.dumps(row) for row in rows) + "\n",
                encoding="utf-8",
            )

            result = export_results_workbook(
                predictions,
                output,
                metadata={"seed": 42, "reproduction_status": "REPRODUCED"},
                paper_targets={"accuracy": 0.5},
                scp_aliases={"NORM": ["normal ecg"], "MI": ["myocardial infarction"]},
                tolerance=0.0,
                scp_fuzzy_threshold=0.91,
            )
            workbook = load_workbook(output, read_only=True)

        self.assertEqual(result["status"], "NOT REPRODUCED")
        self.assertEqual(result["reproduction_status"], "NOT REPRODUCED")
        self.assertEqual(result["numeric_target_status"], "WITHIN TOLERANCE")
        self.assertTrue(result["classification_directional_targets_met"])
        self.assertTrue(result["all_directional_targets_met"])
        self.assertEqual(result["scp_fuzzy_threshold"], 0.91)
        self.assertFalse(result["reproduction_checks"]["canonical_paper_targets"])
        self.assertFalse(result["reproduction_checks"]["structured_attestation_present"])
        self.assertEqual(result["metrics"]["diagnosis_accuracy"], 1.0)
        self.assertEqual(result["metrics"]["accuracy"], 0.5)
        self.assertIn("diagnosis_log_loss", result["metrics"])
        self.assertIn("diagnosis_ece_15_bin", result["metrics"])
        self.assertEqual(result["metrics"]["diagnosis_probability_mass_max_error"], 0.0)
        self.assertEqual(workbook["Summary"]["A2"].value, "Reproduction status")
        self.assertEqual(workbook["Summary"]["B2"].value, "NOT REPRODUCED")
        self.assertEqual(workbook["Summary"]["A3"].value, "Numeric target status")
        self.assertEqual(workbook["Summary"]["B3"].value, "WITHIN TOLERANCE")
        summary_values = {
            workbook["Summary"].cell(row=index, column=1).value: workbook["Summary"]
            .cell(row=index, column=2)
            .value
            for index in range(2, workbook["Summary"].max_row + 1)
        }
        self.assertEqual(summary_values["SCP fuzzy threshold"], 0.91)
        self.assertIn("numeric only", workbook["Paper Comparison"]["F1"].value)
        provenance_values = {
            workbook["Provenance"].cell(row=index, column=1).value: workbook["Provenance"]
            .cell(row=index, column=2)
            .value
            for index in range(2, workbook["Provenance"].max_row + 1)
        }
        self.assertEqual(provenance_values["reproduction_status"], "NOT REPRODUCED")
        prediction_headers = [cell.value for cell in workbook["Predictions"][1]]
        diagnosis_correct = prediction_headers.index("diagnosis_correct") + 1
        primary_prediction = prediction_headers.index("primary_prediction") + 1
        primary_correct = prediction_headers.index("primary_correct") + 1
        self.assertTrue(workbook["Predictions"].cell(row=3, column=diagnosis_correct).value)
        self.assertEqual(
            workbook["Predictions"].cell(row=3, column=primary_prediction).value,
            "NORM",
        )
        self.assertFalse(workbook["Predictions"].cell(row=3, column=primary_correct).value)
        self.assertEqual(
            workbook.sheetnames,
            [
                "Summary",
                "Paper Comparison",
                "Primary Per Class",
                "Confusion Matrix",
                "Provenance",
                "Predictions",
            ],
        )

    def test_reproduction_requires_every_structured_attestation(self):
        rows = [
            {
                "expected": "NORM",
                "diagnosis": "NORM",
                "rationale": "Model-predicted diagnosis: NORM.",
            }
        ]
        targets = {
            "accuracy": 1.0,
            "precision_macro": 1.0,
            "recall_macro": 1.0,
            "f1_macro": 1.0,
        }
        incomplete_attestation = {
            "reproduction_attestation": {
                "full_fixed_test_coverage": True,
                "same_task_protocol": True,
                "frozen_hashes": True,
                "exact_paper_artifacts": True,
                "exact_metric_implementation": False,
            }
        }
        with tempfile.TemporaryDirectory() as directory:
            predictions = Path(directory) / "predictions.jsonl"
            predictions.write_text(json.dumps(rows[0]) + "\n", encoding="utf-8")
            result = export_results_workbook(
                predictions,
                Path(directory) / "results.xlsx",
                metadata=incomplete_attestation,
                paper_targets=targets,
                tolerance=0.01,
            )

        self.assertEqual(result["numeric_target_status"], "WITHIN TOLERANCE")
        self.assertEqual(result["reproduction_status"], "NOT REPRODUCED")
        self.assertFalse(result["reproduction_checks"]["canonical_paper_targets"])
        self.assertFalse(result["reproduction_checks"]["attested_exact_metric_implementation"])

    def test_excel_report_rejects_invalid_scp_fuzzy_threshold(self):
        with tempfile.TemporaryDirectory() as directory:
            predictions = Path(directory) / "predictions.jsonl"
            predictions.write_text(
                json.dumps({"expected": "NORM", "diagnosis": "NORM"}) + "\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "SCP fuzzy threshold"):
                export_results_workbook(
                    predictions,
                    Path(directory) / "results.xlsx",
                    scp_aliases={"NORM": ["normal ECG"]},
                    scp_fuzzy_threshold=1.1,
                )

    def test_crc_is_noncomparable_when_causal_attribution_is_incomplete(self):
        rows = [
            {
                "expected": "NORM",
                "diagnosis": "NORM",
                "query": "diagnose",
                "rationale": "No causal attribution.",
                "retrieved_facts": [],
                "causal_drivers": [],
            },
            {
                "expected": "MI",
                "diagnosis": "MI",
                "query": "diagnose",
                "rationale": "z0 supports the result.",
                "retrieved_facts": [],
                "causal_drivers": [
                    {
                        "name": "z0",
                        "state": 1,
                        "importance": 0.5,
                        "effect": 0.5,
                        "alternative_state": 0,
                        "path": ["z0", "diagnosis"],
                    }
                ],
            },
        ]
        with tempfile.TemporaryDirectory() as directory:
            predictions = Path(directory) / "predictions.jsonl"
            output = Path(directory) / "results.xlsx"
            predictions.write_text(
                "\n".join(json.dumps(row) for row in rows) + "\n",
                encoding="utf-8",
            )
            result = export_results_workbook(
                predictions,
                output,
                metadata={
                    "reproduction_attestation": {"exact_metric_implementation": True}
                },
                paper_targets={"crc": 0.5},
                tolerance=0.0,
            )
            workbook = load_workbook(output, read_only=True)

        self.assertEqual(result["metrics"]["crc"], 0.5)
        self.assertEqual(result["metrics"]["causal_attribution_availability"], 0.5)
        self.assertEqual(result["numeric_target_status"], "OUTSIDE TOLERANCE")
        self.assertTrue(result["all_directional_targets_met"])
        self.assertIn("NO - causal attribution", workbook["Paper Comparison"]["E2"].value)
        self.assertEqual(workbook["Paper Comparison"]["F2"].value, "NO")

    def test_crc_is_noncomparable_when_attribution_availability_is_missing(self):
        row = {
            "expected": "NORM",
            "diagnosis": "NORM",
            "crc": 0.86,
        }
        with tempfile.TemporaryDirectory() as directory:
            predictions = Path(directory) / "predictions.jsonl"
            output = Path(directory) / "results.xlsx"
            predictions.write_text(json.dumps(row) + "\n", encoding="utf-8")
            result = export_results_workbook(
                predictions,
                output,
                metadata={
                    "reproduction_attestation": {"exact_metric_implementation": True}
                },
                paper_targets={"crc": 0.86},
                tolerance=0.0,
            )
            workbook = load_workbook(output, read_only=True)

        self.assertEqual(result["numeric_target_status"], "OUTSIDE TOLERANCE")
        self.assertTrue(result["all_directional_targets_met"])
        self.assertIn("NO - causal attribution", workbook["Paper Comparison"]["E2"].value)
        self.assertEqual(
            workbook["Paper Comparison"]["G2"].value,
            "YES - numeric only; not paper-comparable",
        )

    def test_local_report_metrics_are_not_paper_comparable_without_attestation(self):
        row = {
            "expected": "NORM",
            "diagnosis": "NORM",
            "groundedness": 0.80,
            "hallucination_score": 0.05,
        }
        with tempfile.TemporaryDirectory() as directory:
            predictions = Path(directory) / "predictions.jsonl"
            output = Path(directory) / "results.xlsx"
            predictions.write_text(json.dumps(row) + "\n", encoding="utf-8")
            result = export_results_workbook(
                predictions,
                output,
                paper_targets={"groundedness": 0.67, "hallucination_rate": 0.08},
                tolerance=0.2,
            )
            workbook = load_workbook(output, read_only=True)

        self.assertEqual(result["numeric_target_status"], "OUTSIDE TOLERANCE")
        for row_number in (2, 3):
            self.assertIn(
                "paper-exact metric implementation is not attested",
                workbook["Paper Comparison"].cell(row=row_number, column=5).value,
            )
            self.assertEqual(
                workbook["Paper Comparison"].cell(row=row_number, column=7).value,
                "YES - numeric only; not paper-comparable",
            )


if __name__ == "__main__":
    unittest.main()
