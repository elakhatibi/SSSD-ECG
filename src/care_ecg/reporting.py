from __future__ import annotations

import json
import platform
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

import numpy as np

from .metrics import (
    SCPTextMapper,
    classification_metrics,
    grounding_metrics,
    scp_text_metrics,
    verifier_diagnostics,
)
from .types import CausalDriver, RetrievedFact

PAPER_LLAMA_PTBXL_TARGETS = {
    "accuracy": 0.70,
    "precision_macro": 0.68,
    "recall_macro": 0.72,
    "f1_macro": 0.70,
    "crc": 0.86,
    "groundedness": 0.67,
    "context_relevance": 0.71,
    "hallucination_rate": 0.08,
    "semantic_relevance": 0.77,
}

REPRODUCTION_ATTESTATION_FIELDS = (
    "full_fixed_test_coverage",
    "same_task_protocol",
    "frozen_hashes",
    "exact_paper_artifacts",
    "exact_metric_implementation",
)
CANONICAL_MATCH_TOLERANCE = 0.01
PAPER_REPORT_METRICS = frozenset(
    {
        "crc",
        "groundedness",
        "context_relevance",
        "hallucination_rate",
        "semantic_relevance",
    }
)
LOWER_IS_BETTER_METRICS = frozenset({"hallucination_rate"})


def load_prediction_jsonl(path: str | Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with Path(path).open("r", encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ValueError(f"Invalid prediction JSON at line {line_number}") from exc
            if not isinstance(row, dict):
                raise TypeError(f"Prediction line {line_number} must be a JSON object")
            rows.append(row)
    if not rows:
        raise ValueError("Prediction JSONL is empty")
    return rows


def _classification_summary(
    rows: Sequence[Mapping[str, Any]],
) -> tuple[dict[str, float], list[dict[str, float | str]], list[str], np.ndarray]:
    if not all("expected" in row and "diagnosis" in row for row in rows):
        raise ValueError("Every prediction must contain expected and diagnosis")
    expected = [str(row["expected"]) for row in rows]
    predicted = [str(row["diagnosis"]) for row in rows]
    labels = sorted(set(expected) | set(predicted))
    matrix = np.zeros((len(labels), len(labels)), dtype=np.int64)
    lookup = {label: index for index, label in enumerate(labels)}
    for actual, guess in zip(expected, predicted):
        matrix[lookup[actual], lookup[guess]] += 1

    per_class: list[dict[str, float | str]] = []
    for index, label in enumerate(labels):
        tp = int(matrix[index, index])
        fp = int(matrix[:, index].sum() - tp)
        fn = int(matrix[index, :].sum() - tp)
        support = int(matrix[index, :].sum())
        precision = tp / (tp + fp) if tp + fp else 0.0
        recall = tp / (tp + fn) if tp + fn else 0.0
        f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
        per_class.append(
            {
                "label": label,
                "precision": precision,
                "recall": recall,
                "f1": f1,
                "support": float(support),
            }
        )

    aggregate = classification_metrics(expected, predicted)
    if all("used_fallback" in row for row in rows):
        aggregate.update(
            verifier_diagnostics(
                expected,
                predicted,
                [bool(row["used_fallback"]) for row in rows],
            )
        )
    return aggregate, per_class, labels, matrix


def _wilson_interval(successes: int, total: int, z: float = 1.959963984540054) -> tuple[float, float]:
    if total <= 0:
        raise ValueError("Wilson interval requires at least one observation")
    estimate = successes / total
    denominator = 1.0 + z * z / total
    center = (estimate + z * z / (2.0 * total)) / denominator
    radius = (
        z
        * np.sqrt(estimate * (1.0 - estimate) / total + z * z / (4.0 * total * total))
        / denominator
    )
    return float(max(0.0, center - radius)), float(min(1.0, center + radius))


def _probability_summary(rows: Sequence[Mapping[str, Any]]) -> dict[str, float]:
    """Audit confidence quality when every row contains a class-probability mapping."""

    if not rows or not all(isinstance(row.get("probabilities"), Mapping) for row in rows):
        return {}
    labels = sorted(
        {str(row["expected"]) for row in rows}
        | {str(label) for row in rows for label in row["probabilities"]}
    )
    label_index = {label: index for index, label in enumerate(labels)}
    confidences: list[float] = []
    correct: list[bool] = []
    mass_errors: list[float] = []
    negative_log_likelihood: list[float] = []
    brier_scores: list[float] = []
    for row in rows:
        expected = str(row["expected"])
        diagnosis = str(row["diagnosis"])
        if expected not in label_index or diagnosis not in label_index:
            raise ValueError("Expected and diagnosed labels must be represented in probabilities")
        vector = np.asarray(
            [float(row["probabilities"].get(label, 0.0)) for label in labels], dtype=float
        )
        if not np.all(np.isfinite(vector)) or np.any(vector < 0) or vector.sum() <= 0:
            raise ValueError("Probabilities must be finite, non-negative, and have positive mass")
        mass_errors.append(abs(float(vector.sum()) - 1.0))
        vector = vector / vector.sum()
        truth = label_index[expected]
        one_hot = np.zeros(len(labels), dtype=float)
        one_hot[truth] = 1.0
        negative_log_likelihood.append(-float(np.log(np.clip(vector[truth], 1e-15, 1.0))))
        brier_scores.append(float(np.sum((vector - one_hot) ** 2)))
        confidences.append(float(vector[label_index[diagnosis]]))
        correct.append(expected == diagnosis)

    confidence = np.asarray(confidences)
    correctness = np.asarray(correct, dtype=float)
    calibration_error = 0.0
    for lower, upper in zip(np.linspace(0.0, 1.0, 16)[:-1], np.linspace(0.0, 1.0, 16)[1:]):
        selected = (confidence > lower) & (confidence <= upper)
        if lower == 0.0:
            selected |= confidence == 0.0
        if np.any(selected):
            calibration_error += float(selected.mean()) * abs(
                float(correctness[selected].mean()) - float(confidence[selected].mean())
            )
    keep = max(1, int(np.ceil(0.8 * len(rows))))
    most_confident = np.argsort(-confidence, kind="stable")[:keep]
    return {
        "diagnosis_log_loss": float(np.mean(negative_log_likelihood)),
        "diagnosis_multiclass_brier": float(np.mean(brier_scores)),
        "diagnosis_ece_15_bin": float(calibration_error),
        "diagnosis_mean_confidence": float(confidence.mean()),
        "diagnosis_confidence_gap": float(confidence.mean() - correctness.mean()),
        "diagnosis_accuracy_at_80pct_coverage": float(correctness[most_confident].mean()),
        "diagnosis_probability_mass_max_error": float(max(mass_errors)),
    }


def _mean_recorded_metrics(rows: Sequence[Mapping[str, Any]]) -> dict[str, float]:
    aliases = {
        "crc": "crc",
        "groundedness": "groundedness",
        "context_relevance": "context_relevance",
        "hallucination_score": "hallucination_rate",
        "unsupported_claim_score": "unsupported_claim_rate",
        "semantic_relevance": "semantic_relevance",
    }
    metrics: dict[str, float] = {}
    for source, destination in aliases.items():
        values = [float(row[source]) for row in rows if row.get(source) is not None]
        if values:
            metrics[destination] = float(np.mean(values))
    grounded = []
    for row in rows:
        if not all(key in row for key in ("query", "rationale", "retrieved_facts")):
            continue
        grounded.append(
            grounding_metrics(
                str(row["query"]),
                str(row["rationale"]),
                [RetrievedFact(**value) for value in row["retrieved_facts"]],
                [CausalDriver(**value) for value in row.get("causal_drivers", [])],
            )
        )
    if grounded:
        for name in grounded[0]:
            metrics[name] = float(np.mean([record[name] for record in grounded]))
    return metrics


def _flatten(prefix: str, value: Any) -> list[tuple[str, Any]]:
    if isinstance(value, Mapping):
        rows: list[tuple[str, Any]] = []
        for key in sorted(value, key=str):
            name = f"{prefix}.{key}" if prefix else str(key)
            rows.extend(_flatten(name, value[key]))
        return rows
    if isinstance(value, (list, tuple)):
        return [(prefix, json.dumps(value, sort_keys=True))]
    return [(prefix, value)]


def _write_table(sheet, headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        sheet.append(list(row))
    sheet.freeze_panes = "A2"
    if sheet.max_row > 1:
        sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 60)


def _reproduction_checks(
    metadata: Mapping[str, Any],
    targets: Mapping[str, float],
    tolerance: float,
    numeric_targets_matched: bool,
) -> dict[str, bool]:
    attestation = metadata.get("reproduction_attestation")
    attestation_values = attestation if isinstance(attestation, Mapping) else {}
    checks = {
        "numeric_targets_within_tolerance": numeric_targets_matched,
        "canonical_paper_targets": dict(targets) == PAPER_LLAMA_PTBXL_TARGETS,
        "canonical_absolute_tolerance": bool(
            np.isclose(tolerance, CANONICAL_MATCH_TOLERANCE, rtol=0.0, atol=1e-12)
        ),
        "structured_attestation_present": isinstance(attestation, Mapping),
    }
    checks.update(
        {
            f"attested_{field}": attestation_values.get(field) is True
            for field in REPRODUCTION_ATTESTATION_FIELDS
        }
    )
    return checks


def export_results_workbook(
    predictions: str | Path,
    output: str | Path,
    metadata: Mapping[str, Any] | None = None,
    paper_targets: Mapping[str, float] | None = None,
    scp_aliases: Mapping[str, Sequence[str]] | None = None,
    tolerance: float = 0.01,
    scp_fuzzy_threshold: float = 0.84,
) -> dict[str, Any]:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise ImportError("Excel export requires `pip install 'care-ecg[reports]'`") from exc

    if tolerance < 0:
        raise ValueError("Match tolerance must be non-negative")
    if not 0 <= scp_fuzzy_threshold <= 1:
        raise ValueError("SCP fuzzy threshold must be in [0, 1]")
    rows = load_prediction_jsonl(predictions)
    diagnosis_metrics, per_class, labels, confusion = _classification_summary(rows)
    metrics = {f"diagnosis_{name}": value for name, value in diagnosis_metrics.items()}
    direct_correct = sum(str(row["expected"]) == str(row["diagnosis"]) for row in rows)
    direct_low, direct_high = _wilson_interval(direct_correct, len(rows))
    metrics.update(
        {
            "diagnosis_accuracy_ci95_low": direct_low,
            "diagnosis_accuracy_ci95_high": direct_high,
            **_probability_summary(rows),
        }
    )
    evaluation_mode = "direct diagnosis"
    primary_predictions = [str(row["diagnosis"]) for row in rows]
    if scp_aliases is not None:
        mapper = SCPTextMapper(scp_aliases, scp_fuzzy_threshold)
        expected = [str(row["expected"]) for row in rows]
        texts = [str(row.get("rationale", "")) for row in rows]
        text_metrics = scp_text_metrics(
            expected,
            texts,
            mapper,
        )
        metrics.update(text_metrics)
        for name in ("accuracy", "precision_macro", "recall_macro", "f1_macro"):
            metrics[name] = text_metrics[f"scp_first_{name}"]
        mapped = [mapper.map(text) for text in texts]
        primary_predictions = [codes[0] if codes else "__NO_SCP_MATCH__" for codes in mapped]
        primary_rows = [
            {
                "expected": actual,
                "diagnosis": primary,
            }
            for actual, primary in zip(expected, primary_predictions)
        ]
        _, per_class, labels, confusion = _classification_summary(primary_rows)
        evaluation_mode = "first-mention Text-to-SCP"
    else:
        metrics.update(diagnosis_metrics)
    primary_correct = sum(
        str(row["expected"]) == prediction for row, prediction in zip(rows, primary_predictions)
    )
    primary_low, primary_high = _wilson_interval(primary_correct, len(rows))
    metrics["accuracy_ci95_low"] = primary_low
    metrics["accuracy_ci95_high"] = primary_high
    metrics.update(_mean_recorded_metrics(rows))
    attribution_availability = metrics.get("causal_attribution_availability")
    fused_attribution_availability = metrics.get("fused_causal_attribution_availability")
    crc_attribution_complete = bool(
        attribution_availability is not None
        and fused_attribution_availability is not None
        and np.isclose(attribution_availability, 1.0, rtol=0.0, atol=1e-12)
        and np.isclose(fused_attribution_availability, 1.0, rtol=0.0, atol=1e-12)
    )
    metadata_values = dict(metadata or {})
    attestation = metadata_values.get("reproduction_attestation")
    exact_metric_implementation = bool(
        isinstance(attestation, Mapping)
        and attestation.get("exact_metric_implementation") is True
    )
    crc_comparable = exact_metric_implementation and crc_attribution_complete
    targets = dict(paper_targets or PAPER_LLAMA_PTBXL_TARGETS)
    comparisons = []
    all_numeric_targets_matched = True
    all_directional_targets_met = True
    classification_directional_targets_met = True
    for name, target in targets.items():
        measured = metrics.get(name)
        delta = measured - target if measured is not None else None
        comparable = name not in PAPER_REPORT_METRICS or exact_metric_implementation
        comparability = "YES"
        if name in PAPER_REPORT_METRICS and not exact_metric_implementation:
            comparability = "NO - paper-exact metric implementation is not attested"
        elif name == "crc" and not crc_attribution_complete:
            comparable = False
            comparability = (
                "NO - causal attribution is incomplete for the reported fused posterior"
            )
        matched = measured is not None and comparable and abs(delta) <= tolerance
        all_numeric_targets_matched = all_numeric_targets_matched and matched
        if measured is None:
            directional_status = "N/A - metric unavailable"
            directional_target_met = False
        else:
            directional_target_met = (
                measured <= target if name in LOWER_IS_BETTER_METRICS else measured >= target
            )
            directional_status = "YES" if directional_target_met else "NO"
            if not comparable:
                directional_status += " - numeric only; not paper-comparable"
        all_directional_targets_met = (
            all_directional_targets_met and directional_target_met
        )
        if name not in PAPER_REPORT_METRICS:
            classification_directional_targets_met = (
                classification_directional_targets_met and directional_target_met
            )
        comparisons.append(
            (
                name,
                target,
                measured,
                delta,
                comparability,
                "YES" if matched else "NO",
                directional_status,
            )
        )
    numeric_target_status = (
        "WITHIN TOLERANCE" if all_numeric_targets_matched else "OUTSIDE TOLERANCE"
    )
    reproduction_checks = _reproduction_checks(
        metadata_values,
        targets,
        tolerance,
        all_numeric_targets_matched,
    )
    reproduction_status = "REPRODUCED" if all(reproduction_checks.values()) else "NOT REPRODUCED"

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary_rows = [
        ("Reproduction status", reproduction_status),
        ("Numeric target status", numeric_target_status),
        (
            "Classification directional targets",
            "MET OR EXCEEDED" if classification_directional_targets_met else "NOT ALL MET",
        ),
        (
            "All directional numeric targets",
            "MET OR EXCEEDED" if all_directional_targets_met else "NOT ALL MET",
        ),
        (
            "Interpretation",
            "Numeric proximity alone does not establish reproduction.",
        ),
        ("Prediction records", len(rows)),
        ("Paper target", "CARE-ECG LLaMA 3-8B on PTB-XL"),
        ("Primary classification", evaluation_mode),
        ("Absolute tolerance", tolerance),
        ("SCP fuzzy threshold", scp_fuzzy_threshold if scp_aliases is not None else "N/A"),
        (
            "CRC comparability",
            "COMPARABLE"
            if crc_comparable
            else (
                "NOT COMPARABLE: paper-exact metric implementation is not attested"
                if not exact_metric_implementation
                else "NOT COMPARABLE: complete fused-posterior attribution is unavailable"
            ),
        ),
    ]
    summary_rows.extend(
        (f"Reproduction check: {name}", "YES" if passed else "NO")
        for name, passed in reproduction_checks.items()
    )
    summary_rows.extend((name, value) for name, value in sorted(metrics.items()))
    _write_table(summary, ("Field", "Value"), summary_rows)
    status_cell = summary["B2"]
    status_cell.font = Font(bold=True, color="FFFFFF")
    status_cell.fill = PatternFill(
        "solid", fgColor="548235" if reproduction_status == "REPRODUCED" else "C00000"
    )

    comparison = workbook.create_sheet("Paper Comparison")
    _write_table(
        comparison,
        (
            "Metric",
            "Paper target",
            "Measured",
            "Delta",
            "Comparable to paper target",
            f"Within absolute tolerance {tolerance:g} (numeric only)",
            "Meets directional target (numeric only when non-comparable)",
        ),
        comparisons,
    )

    class_sheet = workbook.create_sheet("Primary Per Class")
    _write_table(
        class_sheet,
        ("Label", "Precision", "Recall", "F1", "Support"),
        [
            (row["label"], row["precision"], row["recall"], row["f1"], int(row["support"]))
            for row in per_class
        ],
    )

    confusion_sheet = workbook.create_sheet("Confusion Matrix")
    _write_table(
        confusion_sheet,
        ("Actual / Predicted", *labels),
        [
            (label, *[int(value) for value in confusion[index]])
            for index, label in enumerate(labels)
        ],
    )

    provenance = {
        "predictions_path": str(Path(predictions).resolve()),
        "python": platform.python_version(),
        "platform": platform.platform(),
        **metadata_values,
        "scp_fuzzy_threshold": scp_fuzzy_threshold if scp_aliases is not None else None,
        "numeric_target_status": numeric_target_status,
        "reproduction_status": reproduction_status,
        "reproduction_checks": reproduction_checks,
    }
    provenance_sheet = workbook.create_sheet("Provenance")
    _write_table(provenance_sheet, ("Field", "Value"), _flatten("", provenance))

    prediction_sheet = workbook.create_sheet("Predictions")
    prediction_headers = [
        "record_id",
        "expected",
        "diagnosis",
        "diagnosis_correct",
        "primary_prediction",
        "primary_correct",
        "used_fallback",
        "abstained",
        "hallucination_score",
        "unsupported_claim_score",
        "top_probability",
        "probability_margin",
        "causal_driver_count",
        "causal_attribution_scope",
        "retrieved_fact_count",
        "probabilities",
    ]
    prediction_values = []
    for index, (row, primary_prediction) in enumerate(zip(rows, primary_predictions)):
        row_probabilities = row.get("probabilities", {})
        ranked_probabilities = (
            sorted((float(value) for value in row_probabilities.values()), reverse=True)
            if isinstance(row_probabilities, Mapping)
            else []
        )
        top_probability = ranked_probabilities[0] if ranked_probabilities else None
        probability_margin = (
            ranked_probabilities[0] - ranked_probabilities[1]
            if len(ranked_probabilities) >= 2
            else None
        )
        row_metadata = row.get("metadata", {})
        attribution = (
            row_metadata.get("causal_attribution", {})
            if isinstance(row_metadata, Mapping)
            else {}
        )
        if not isinstance(attribution, Mapping):
            attribution = {}
        prediction_values.append(
            (
                row.get("record_id", index),
                row["expected"],
                row["diagnosis"],
                str(row["expected"]) == str(row["diagnosis"]),
                primary_prediction,
                str(row["expected"]) == primary_prediction,
                row.get("used_fallback"),
                row.get("abstained"),
                row.get("hallucination_score"),
                row.get("unsupported_claim_score"),
                top_probability,
                probability_margin,
                len(row.get("causal_drivers", [])),
                attribution.get("scope"),
                len(row.get("retrieved_facts", [])),
                json.dumps(row.get("probabilities", {}), sort_keys=True),
            )
        )
    _write_table(prediction_sheet, prediction_headers, prediction_values)

    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return {
        "status": reproduction_status,
        "reproduction_status": reproduction_status,
        "numeric_target_status": numeric_target_status,
        "classification_directional_targets_met": classification_directional_targets_met,
        "all_directional_targets_met": all_directional_targets_met,
        "scp_fuzzy_threshold": scp_fuzzy_threshold if scp_aliases is not None else None,
        "reproduction_checks": reproduction_checks,
        "metrics": metrics,
        "path": str(output_path.resolve()),
    }
