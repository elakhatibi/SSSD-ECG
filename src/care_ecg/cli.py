from __future__ import annotations

import argparse
import hashlib
import json
import sys
from collections.abc import Mapping, Sequence
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import numpy as np
import yaml

from . import __version__
from .config import CAREConfig
from .data import (
    PTBXL_DIAGNOSTIC_SUPERCLASSES,
    ECGArrayBundle,
    iter_ptbxl_diagnostic_superclasses,
    load_npz,
    load_ptbxl_diagnostic_aliases,
    preprocess_ecg,
    records_to_arrays,
)
from .expert import load_aligned_expert_probabilities
from .knowledge_graph import ClinicalKnowledgeGraph, build_ptbxl_knowledge_graph
from .metrics import (
    SCPTextMapper,
    classification_metrics,
    grounding_metrics,
    scp_text_metrics,
    verifier_diagnostics,
)
from .pipeline import CAREECGPipeline, bundled_demo_documents
from .reporting import PAPER_LLAMA_PTBXL_TARGETS, export_results_workbook
from .retrieval import load_jsonl_documents

SECURITY_CONFIG_SCHEMA = "care-ecg-security-config.v1"
SECURITY_RESULTS_SCHEMA = "care-ecg-security-results.v2"
SECURITY_LOCK_SCHEMA = "care-ecg-security-validation-lock.v2"
SECURITY_ATTACK_NAMES = (
    "f0",
    "f1_pgd",
    "f1_cf",
    "f2_nes",
    "f2_simba",
    "f2_square",
    "f3_k1",
    "f3_k2",
    "f3_k3",
)
SECURITY_CONSTRAINT_NAMES = ("s_inf", "s2", "s_spec", "s_phys")


def _load_config(path: str | None) -> CAREConfig:
    return CAREConfig.from_yaml(path) if path else CAREConfig()


def _sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _directory_identity(path: str | Path, *, suffix: str | None = None) -> dict[str, Any]:
    root = Path(path).resolve()
    if not root.is_dir():
        raise FileNotFoundError(f"Identity directory does not exist: {root}")
    files = {
        str(item.relative_to(root)): _sha256_file(item)
        for item in sorted(root.rglob("*"))
        if item.is_file() and (suffix is None or item.suffix == suffix)
    }
    if not files:
        raise ValueError(f"Identity directory has no matching files: {root}")
    digest = hashlib.sha256(
        json.dumps(files, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    return {"sha256": digest, "files": files}


def _security_code_identity() -> dict[str, Any]:
    return _directory_identity(Path(__file__).resolve().parent, suffix=".py")


def _prepare_signals(
    signals: np.ndarray,
    config: CAREConfig,
    source_rate: int | None,
    preprocessed: bool = False,
) -> np.ndarray:
    if signals.ndim != 3:
        raise ValueError("Signals must have shape [records, leads, samples]")
    expected = (config.data.leads, config.data.samples)
    if preprocessed:
        if signals.shape[1:] != expected:
            raise ValueError(f"Preprocessed signals must have shape [records, {expected}]")
        if not np.all(np.isfinite(signals)):
            raise ValueError("Preprocessed signals must be finite")
        return np.ascontiguousarray(signals, dtype=np.float32)
    if source_rate is None:
        if signals.shape[1:] not in {expected, expected[::-1]}:
            raise ValueError(
                f"Signals have shape {signals.shape[1:]}, expected {expected}; "
                "provide --source-rate"
            )
        source_rate = config.data.sample_rate
    normalization = "none" if config.data.normalization == "global" else config.data.normalization
    return np.stack(
        [
            preprocess_ecg(
                signal,
                source_rate,
                config.data.sample_rate,
                config.data.seconds,
                config.data.leads,
                normalization,
            )
            for signal in signals
        ]
    )


def _train(args: argparse.Namespace) -> int:
    config = _load_config(args.config)
    bundle = load_npz(args.dataset, require_labels=True)
    signals = _prepare_signals(bundle.signals, config, args.source_rate, args.preprocessed)
    raw_labels = np.asarray(bundle.labels)
    label_lookup = None
    if args.label_names:
        label_names = [part.strip() for part in args.label_names.split(",") if part.strip()]
        labels = raw_labels.astype(np.int64)
    else:
        unique = sorted(set(raw_labels.tolist()), key=str)
        label_names = [str(value) for value in unique]
        label_lookup = {value: i for i, value in enumerate(unique)}
        labels = np.asarray([label_lookup[value] for value in raw_labels.tolist()], dtype=np.int64)
    documents = load_jsonl_documents(args.knowledge) if args.knowledge else bundled_demo_documents()
    knowledge_graph = None
    if args.knowledge_graph:
        knowledge_graph = ClinicalKnowledgeGraph.load(args.knowledge_graph)
        documents.extend(knowledge_graph.retrieval_documents())
    validation_signals = None
    validation_labels = None
    if args.validation_dataset:
        validation_bundle = load_npz(args.validation_dataset, require_labels=True)
        validation_signals = _prepare_signals(
            validation_bundle.signals,
            config,
            args.validation_source_rate
            if args.validation_source_rate is not None
            else args.source_rate,
            args.preprocessed,
        )
        raw_validation_labels = np.asarray(validation_bundle.labels)
        if label_lookup is None:
            validation_labels = raw_validation_labels.astype(np.int64)
        else:
            unknown = set(raw_validation_labels.tolist()) - set(label_lookup)
            if unknown:
                raise ValueError(f"Validation labels were absent from training: {sorted(unknown)}")
            validation_labels = np.asarray(
                [label_lookup[value] for value in raw_validation_labels.tolist()],
                dtype=np.int64,
            )
    prior_edges = []
    if args.prior_edges:
        prior_edges = [tuple(edge) for edge in json.loads(Path(args.prior_edges).read_text())]
    if knowledge_graph is not None:
        variable_order = [f"z{index}" for index in range(config.model.latent_dim)] + ["diagnosis"]
        prior_edges.extend(knowledge_graph.causal_prior_edges(variable_order))
    prior_edges = list(dict.fromkeys(prior_edges))

    def progress(epoch, train, validation):
        val = f", val={validation['loss']:.5f}" if validation else ""
        if "macro_f1" in validation:
            val += f", val_macro_f1={validation['macro_f1']:.5f}"
        print(f"epoch {epoch:03d}: train={train['loss']:.5f}{val}", file=sys.stderr)

    pipeline, history = CAREECGPipeline.train(
        signals=signals,
        labels=labels,
        label_names=label_names,
        documents=documents,
        config=config,
        clinical_prior_edges=prior_edges,
        validation_fraction=args.validation_fraction,
        progress=progress,
        validation_signals=validation_signals,
        validation_labels=validation_labels,
    )
    pipeline.save(args.output)
    if knowledge_graph is not None:
        knowledge_graph.save(Path(args.output) / "knowledge_graph.json")
    from .training import CausalVAETrainer

    CausalVAETrainer(pipeline.model, config).save_history(
        history, Path(args.output) / "training_history.json"
    )
    print(
        "latent activity: "
        f"{history.diagnostics['active_units']}/"
        f"{history.diagnostics['latent_dimensions']} units",
        file=sys.stderr,
    )
    print(f"Saved CARE-ECG artifacts to {Path(args.output).resolve()}")
    return 0


def _save_npz_bundle(path: Path, bundle: ECGArrayBundle, mask: np.ndarray) -> None:
    values: dict[str, np.ndarray] = {
        "signals": bundle.signals[mask],
        "labels": np.asarray(bundle.labels)[mask],
        "folds": np.asarray(bundle.folds)[mask],
    }
    for name in ("patient_ids", "timestamps", "record_ids"):
        value = getattr(bundle, name)
        if value is not None:
            values[name] = np.asarray(value)[mask]
    np.savez_compressed(path, **values)


def _prepare_ptbxl(args: argparse.Namespace) -> int:
    output = Path(args.output)
    output.mkdir(parents=True, exist_ok=True)
    bundle = records_to_arrays(
        iter_ptbxl_diagnostic_superclasses(
            args.root,
            sampling_rate=args.sampling_rate,
            min_likelihood=args.min_likelihood,
        ),
        PTBXL_DIAGNOSTIC_SUPERCLASSES,
        target_rate=args.sampling_rate,
        seconds=args.seconds,
        normalization=args.normalization,
    )
    if bundle.folds is None:
        raise RuntimeError("PTB-XL records did not contain strat_fold metadata")
    folds = np.asarray(bundle.folds)
    split_masks = {
        "train": np.isin(folds, np.arange(1, 9)),
        "validation": folds == 9,
        "test": folds == 10,
    }
    counts: dict[str, int] = {}
    for name, mask in split_masks.items():
        counts[name] = int(mask.sum())
        if not counts[name]:
            raise ValueError(f"PTB-XL {name} split is empty")
        _save_npz_bundle(output / f"{name}.npz", bundle, mask)
    manifest = {
        "dataset": "PTB-XL",
        "sampling_rate": args.sampling_rate,
        "seconds": args.seconds,
        "normalization": args.normalization,
        "labels": list(PTBXL_DIAGNOSTIC_SUPERCLASSES),
        "label_policy": "highest-likelihood diagnostic superclass; ties follow label order",
        "folds": {"train": list(range(1, 9)), "validation": [9], "test": [10]},
        "records": counts,
    }
    knowledge_graph = build_ptbxl_knowledge_graph(args.root)
    knowledge_graph.save(output / "knowledge_graph.json")
    manifest["knowledge_graph"] = {
        "file": "knowledge_graph.json",
        "nodes": len(knowledge_graph.nodes),
        "edges": len(knowledge_graph.edges),
        "causal_priors": 0,
    }
    (output / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    (output / "scp_aliases.json").write_text(
        json.dumps(load_ptbxl_diagnostic_aliases(args.root), indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(manifest, indent=2))
    return 0


def _infer(args: argparse.Namespace) -> int:
    config_override = None
    requested_backend = "openai" if args.openai else args.generator
    if requested_backend:
        config_override = CAREConfig.from_yaml(Path(args.artifacts) / "config.yaml")
        config_override.generation.backend = requested_backend
        if args.model:
            config_override.generation.model = args.model
    pipeline = CAREECGPipeline.load(args.artifacts, config_override)
    signal = np.load(args.signal, allow_pickle=False)
    history_entries = []
    if args.history:
        history = load_npz(args.history)
        history_entries = pipeline.make_history_entries(
            history.signals,
            history.record_ids,
            history.patient_ids,
            history.timestamps,
            args.history_source_rate,
        )
    output = pipeline.run(
        signal,
        args.query,
        source_rate=args.source_rate,
        history=history_entries,
        patient_id=args.patient_id,
        counterfactual_target=args.counterfactual_target,
    )
    serialized = json.dumps(output.to_dict(), indent=2)
    if args.output:
        Path(args.output).write_text(serialized + "\n", encoding="utf-8")
    else:
        print(serialized)
    return 0


def _predict(args: argparse.Namespace) -> int:
    config_override = CAREConfig.from_yaml(Path(args.artifacts) / "config.yaml")
    if args.generator:
        config_override.generation.backend = args.generator
    if args.model:
        config_override.generation.model = args.model
    pipeline = CAREECGPipeline.load(args.artifacts, config_override)
    bundle = load_npz(args.dataset, require_labels=True)
    signals = _prepare_signals(
        bundle.signals,
        pipeline.config,
        args.source_rate,
        args.preprocessed,
    )
    count = len(signals) if args.limit is None else min(len(signals), args.limit)
    if count < 1:
        raise ValueError("Prediction dataset is empty")
    expert_probabilities = None
    expert_artifact = None
    expert_path = getattr(args, "expert_probabilities", None)
    expert_weight = getattr(args, "expert_weight", None)
    expert_source = getattr(args, "expert_source", None)
    expert_options = (expert_path is not None, expert_weight is not None, expert_source is not None)
    if any(expert_options) and not all(expert_options):
        raise ValueError(
            "--expert-probabilities, --expert-weight, and --expert-source must be used together"
        )
    if all(expert_options):
        expert_probabilities = load_aligned_expert_probabilities(
            expert_path,
            bundle.record_ids,
            pipeline.causal_model.label_names,
        )
        resolved_expert_path = Path(expert_path).resolve()
        expert_artifact = {
            "path": str(resolved_expert_path),
            "sha256": _sha256_file(resolved_expert_path),
            "size_bytes": resolved_expert_path.stat().st_size,
        }
    latent_values, neural_probabilities = pipeline.encode_batch_with_neural(
        signals[:count], batch_size=args.batch_size
    )
    rows: list[str] = []
    raw_labels = np.asarray(bundle.labels)
    for index in range(count):
        raw_label = (
            raw_labels[index].item() if hasattr(raw_labels[index], "item") else raw_labels[index]
        )
        if isinstance(raw_label, (int, np.integer)):
            label_index = int(raw_label)
            if not 0 <= label_index < len(pipeline.causal_model.label_names):
                raise ValueError(f"Label index {label_index} is outside the trained label set")
            expected = pipeline.causal_model.label_names[label_index]
        else:
            expected = str(raw_label)
        result = pipeline.run_encoded(
            latent_values[index],
            args.query,
            counterfactual_target=args.counterfactual_target,
            neural_probabilities=None
            if neural_probabilities is None
            else neural_probabilities[index],
            external_expert_probabilities=None
            if expert_probabilities is None
            else expert_probabilities[index],
            external_expert_weight=expert_weight,
            external_expert_source=expert_source,
        ).to_dict()
        result["metadata"]["inference"] = {
            "encoder_batch_size": args.batch_size,
            "deterministic_latent_sampling": True,
        }
        if expert_artifact is not None:
            result["metadata"]["diagnosis_fusion"]["external_expert"]["probability_artifact"] = (
                expert_artifact
            )
        result["expected"] = expected
        if bundle.record_ids is not None:
            result["record_id"] = str(bundle.record_ids[index])
        rows.append(json.dumps(result, separators=(",", ":")))
        if (index + 1) % 25 == 0 or index + 1 == count:
            print(f"predicted {index + 1}/{count}", file=sys.stderr)
    Path(args.output).write_text("\n".join(rows) + "\n", encoding="utf-8")
    print(f"Saved {count} predictions to {Path(args.output).resolve()}")
    return 0


def _evaluate(args: argparse.Namespace) -> int:
    rows = [
        json.loads(line) for line in Path(args.predictions).read_text().splitlines() if line.strip()
    ]
    metrics: dict[str, float] = {}
    if all("expected" in row and "diagnosis" in row for row in rows):
        metrics.update(
            classification_metrics(
                [row["expected"] for row in rows], [row["diagnosis"] for row in rows]
            )
        )
    if all("expected" in row and "diagnosis" in row and "used_fallback" in row for row in rows):
        metrics.update(
            verifier_diagnostics(
                [row["expected"] for row in rows],
                [row["diagnosis"] for row in rows],
                [bool(row["used_fallback"]) for row in rows],
            )
        )
    if args.scp_aliases:
        aliases = json.loads(Path(args.scp_aliases).read_text(encoding="utf-8"))
        if not isinstance(aliases, dict):
            raise ValueError("SCP aliases must be a JSON object mapping codes to aliases")
        if not all(args.expected_field in row and args.text_field in row for row in rows):
            raise ValueError(
                f"Every row must contain {args.expected_field!r} and {args.text_field!r}"
            )
        metrics.update(
            scp_text_metrics(
                [str(row[args.expected_field]) for row in rows],
                [str(row[args.text_field]) for row in rows],
                SCPTextMapper(aliases, args.scp_fuzzy_threshold),
            )
        )
    grounding = []
    for row in rows:
        if not all(key in row for key in ("query", "rationale", "retrieved_facts")):
            continue
        from .types import CausalDriver, RetrievedFact

        grounding.append(
            grounding_metrics(
                row["query"],
                row["rationale"],
                [RetrievedFact(**value) for value in row["retrieved_facts"]],
                [CausalDriver(**value) for value in row.get("causal_drivers", [])],
            )
        )
    if grounding:
        for name in grounding[0]:
            metrics[name] = float(np.mean([row[name] for row in grounding]))
    print(json.dumps(metrics, indent=2))
    return 0


def _export_results(args: argparse.Namespace) -> int:
    metadata = {}
    if args.metadata:
        metadata = json.loads(Path(args.metadata).read_text(encoding="utf-8"))
        if not isinstance(metadata, dict):
            raise ValueError("Report metadata must be a JSON object")
    targets = PAPER_LLAMA_PTBXL_TARGETS
    if args.paper_targets:
        targets = json.loads(Path(args.paper_targets).read_text(encoding="utf-8"))
        if not isinstance(targets, dict):
            raise ValueError("Paper targets must be a JSON object")
        targets = {str(name): float(value) for name, value in targets.items()}
    aliases = None
    if args.scp_aliases:
        aliases = json.loads(Path(args.scp_aliases).read_text(encoding="utf-8"))
        if not isinstance(aliases, dict):
            raise ValueError("SCP aliases must be a JSON object mapping codes to aliases")
    result = export_results_workbook(
        args.predictions,
        args.output,
        metadata=metadata,
        paper_targets=targets,
        scp_aliases=aliases,
        tolerance=args.tolerance,
        scp_fuzzy_threshold=args.scp_fuzzy_threshold,
    )
    print(json.dumps(result, indent=2))
    return 0


def _load_security_config(path: str | Path) -> dict[str, Any]:
    source = Path(path)
    values = yaml.safe_load(source.read_text(encoding="utf-8"))
    if not isinstance(values, dict):
        raise TypeError("Security configuration must be a YAML mapping")
    if values.get("schema_version") != SECURITY_CONFIG_SCHEMA:
        raise ValueError(
            f"Security configuration schema_version must be {SECURITY_CONFIG_SCHEMA!r}"
        )
    for section in ("protocol", "attacks", "constraints", "admissibility", "export"):
        if not isinstance(values.get(section), dict):
            raise TypeError(f"Security configuration requires a {section!r} mapping")
    return values


def _csv_items(value: str | None, fallback: Sequence[Any], cast=str) -> tuple[Any, ...]:
    if value is None:
        items = tuple(fallback)
    else:
        items = tuple(part.strip() for part in value.split(",") if part.strip())
    if not items:
        raise ValueError("A comma-separated security option must not be empty")
    return tuple(cast(item) for item in items)


def _security_raw_signals(
    signals: np.ndarray,
    pipeline: CAREECGPipeline,
    source_rate: int | None,
) -> np.ndarray:
    values = np.asarray(signals, dtype=np.float32)
    if values.ndim != 3 or not len(values) or not np.isfinite(values).all():
        raise ValueError(
            "Security signals must be a non-empty finite [records, leads, samples] array"
        )
    expected = (pipeline.config.data.leads, pipeline.config.data.samples)
    if source_rate is None:
        if values.shape[1:] != expected:
            raise ValueError(
                f"Security signals must have shape [records, {expected[0]}, {expected[1]}] "
                "or provide --source-rate"
            )
        return np.ascontiguousarray(values)
    return np.stack(
        [
            preprocess_ecg(
                signal,
                source_rate,
                pipeline.config.data.sample_rate,
                pipeline.config.data.seconds,
                pipeline.config.data.leads,
                "none",
            )
            for signal in values
        ]
    ).astype(np.float32, copy=False)


def _security_label_names(labels: np.ndarray, trained_labels: Sequence[str]) -> np.ndarray:
    raw = np.asarray(labels)
    names = tuple(str(label) for label in trained_labels)
    if np.issubdtype(raw.dtype, np.integer):
        indices = raw.astype(np.int64)
        if len(indices) and (indices.min() < 0 or indices.max() >= len(names)):
            raise ValueError("Security dataset label index is outside the trained label set")
        return np.asarray([names[index] for index in indices], dtype=str)
    values = raw.astype(str)
    unknown = sorted(set(values.tolist()) - set(names))
    if unknown:
        raise ValueError(f"Security dataset contains labels absent from the victim: {unknown}")
    return values


def _security_cohort_cases(
    bundle: ECGArrayBundle,
    signals: np.ndarray,
    trained_labels: Sequence[str],
    target_ids: Sequence[str],
    *,
    max_candidates_per_target: int | None = None,
    allow_record_level: bool = False,
) -> list[Any]:
    from .security.contracts import CLINICAL_TARGETS
    from .security.experiments import CohortCase

    if bundle.labels is None:
        raise ValueError("Security evaluation requires labels")
    if max_candidates_per_target is not None and max_candidates_per_target < 1:
        raise ValueError("max_candidates_per_target must be positive")
    unknown_targets = sorted(set(target_ids) - set(CLINICAL_TARGETS))
    if unknown_targets:
        raise ValueError(f"Unknown clinical target IDs: {unknown_targets}")
    label_names = _security_label_names(bundle.labels, trained_labels)
    count = len(signals)
    record_ids = (
        np.asarray(bundle.record_ids).astype(str)
        if bundle.record_ids is not None
        else np.asarray([str(index) for index in range(count)])
    )
    if bundle.patient_ids is None:
        if not allow_record_level:
            raise ValueError(
                "Patient IDs are required for patient-level security evaluation; use "
                "--allow-record-level only for an explicitly record-level exploratory run"
            )
        patient_ids = record_ids
    else:
        patient_ids = np.asarray(bundle.patient_ids).astype(str)
        missing = np.flatnonzero(np.char.strip(patient_ids) == "")
        if len(missing):
            if not allow_record_level:
                raise ValueError("Security dataset contains empty patient IDs")
            patient_ids = patient_ids.copy()
            patient_ids[missing] = record_ids[missing]

    cases: list[CohortCase] = []
    for target_id in target_ids:
        target = CLINICAL_TARGETS[target_id]
        selected_patients: set[str] = set()
        for index, label in enumerate(label_names):
            if label != target.clean_label or patient_ids[index] in selected_patients:
                continue
            metadata: dict[str, Any] = {"source_index": int(index)}
            if bundle.folds is not None:
                metadata["fold"] = int(bundle.folds[index])
            if bundle.timestamps is not None:
                metadata["timestamp"] = float(bundle.timestamps[index])
            cases.append(
                CohortCase(
                    patient_id=str(patient_ids[index]),
                    record_id=str(record_ids[index]),
                    waveform_mv=signals[index],
                    target=target,
                    metadata=metadata,
                )
            )
            selected_patients.add(str(patient_ids[index]))
            if (
                max_candidates_per_target is not None
                and len(selected_patients) >= max_candidates_per_target
            ):
                break
    if not cases:
        raise ValueError("No dataset records match the configured security targets")
    return cases


def _limit_eligible_cases(eligible: Sequence[Any], limit_per_target: int) -> tuple[Any, ...]:
    if limit_per_target < 1:
        raise ValueError("Eligible patient limit must be positive")
    counts: dict[str, int] = {}
    selected: list[Any] = []
    for item in eligible:
        target_id = item.case.target.target_id
        if counts.get(target_id, 0) >= limit_per_target:
            continue
        selected.append(item)
        counts[target_id] = counts.get(target_id, 0) + 1
    return tuple(selected)


def _security_cohort_manifest(cases: Sequence[Any]) -> dict[str, Any]:
    records = [
        {
            "patient_id": str(case.patient_id),
            "record_id": str(case.record_id),
            "target_id": str(case.target.target_id),
            "source_index": case.metadata.get("source_index"),
            "fold": case.metadata.get("fold"),
        }
        for case in cases
    ]
    payload = json.dumps(records, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return {
        "sha256": hashlib.sha256(payload).hexdigest(),
        "records": len(records),
        "patient_target_pairs": len(
            {(record["patient_id"], record["target_id"]) for record in records}
        ),
    }


def _security_constraint(name: str, values: Mapping[str, Any], sample_rate_hz: float):
    from .security.constraints import S2Constraint, SInfConstraint, SPhysConstraint, SSpecConstraint

    if name == "s_inf":
        return SInfConstraint(epsilon_mv=float(values.get("epsilon_mv", 0.01)))
    if name == "s2":
        return S2Constraint(
            epsilon_mv=float(values.get("epsilon_mv", 0.20)),
            per_lead=bool(values.get("per_lead", True)),
        )
    if name == "s_spec":
        max_slew = values.get("max_slew_mv")
        return SSpecConstraint(
            epsilon_mv=float(values.get("epsilon_mv", 0.02)),
            sample_rate_hz=sample_rate_hz,
            cutoff_hz=float(values.get("cutoff_hz", 40.0)),
            max_high_frequency_ratio=float(values.get("max_high_frequency_ratio", 0.05)),
            max_slew_mv=None if max_slew is None else float(max_slew),
        )
    if name == "s_phys":
        kwargs: dict[str, Any] = {
            "sample_rate_hz": sample_rate_hz,
            "nominal_rr_seconds": float(values.get("nominal_rr_seconds", 0.8)),
        }
        if "bounds" in values:
            kwargs["bounds"] = {str(key): float(value) for key, value in values["bounds"].items()}
        return SPhysConstraint(**kwargs)
    raise ValueError(f"Unknown security constraint: {name}")


def _security_attack_factory(
    name: str,
    attack_config: Mapping[str, Any],
    constraint: Any,
    matched_magnitudes: Mapping[int, float] | None = None,
):
    from .security.attacks import (
        F0RandomAttack,
        F1CounterfactualAttack,
        F1PGDAttack,
        F2NESAttack,
        F2SimBAAttack,
        F2SquareAttack,
        F3AdaptiveAttack,
    )
    from .security.contracts import AttackFamily, ConstraintKind, KnowledgeLevel

    if name not in SECURITY_ATTACK_NAMES:
        raise ValueError(f"Unknown security attack: {name}")
    common = dict(attack_config.get("common", {}))
    specific = dict(attack_config.get(name, {}))
    values = {**common, **specific}
    norm = "l2" if constraint.kind is ConstraintKind.S2 else "linf"
    epsilon_mv = float(getattr(constraint, "epsilon_mv", values.get("epsilon_mv", 0.01)))
    step_size_mv = float(values.get("step_size_mv", 0.001))
    steps = int(values.get("steps", 40))
    restarts = int(values.get("restarts", 1))
    random_start = bool(values.get("random_start", True))
    confidence = float(values.get("confidence", 0.0))
    parameterization = constraint if constraint.kind is ConstraintKind.S_PHYS else None

    def factory(*, seed: int, budget: int | None = None, **_: Any):
        gradient = {
            "norm": norm,
            "epsilon_mv": epsilon_mv,
            "step_size_mv": step_size_mv,
            "steps": steps,
            "restarts": restarts,
            "random_start": random_start,
            "confidence": confidence,
            "seed": seed,
        }
        if name == "f0":
            configured_magnitude = values.get("matched_magnitude")
            matched_magnitude = (
                None
                if matched_magnitudes is None or budget is None
                else matched_magnitudes.get(int(budget))
            )
            if matched_magnitude is None:
                matched_magnitude = configured_magnitude
            if matched_magnitude is None:
                raise ValueError(
                    "F0 requires a validation-derived matched magnitude for every budget"
                )
            return F0RandomAttack(
                norm=norm,
                epsilon_mv=epsilon_mv,
                matched_magnitude=float(matched_magnitude),
                seed=seed,
            )
        if name == "f1_pgd":
            return F1PGDAttack(**gradient)
        if name == "f1_cf":
            return F1CounterfactualAttack(
                latent_weight=float(values.get("latent_weight", 1.0)),
                latent_margin=float(values.get("latent_margin", 1e-3)),
                **gradient,
            )
        if name == "f2_nes":
            return F2NESAttack(
                sigma=float(values.get("sigma", 0.001)),
                learning_rate=float(values.get("learning_rate", 0.002)),
                samples_per_step=int(values.get("samples_per_step", 4)),
                epsilon_mv=epsilon_mv,
                norm=norm,
                seed=seed,
                parameterization=parameterization,
            )
        if name == "f2_simba":
            return F2SimBAAttack(
                step_size_mv=step_size_mv,
                epsilon_mv=epsilon_mv,
                norm=norm,
                seed=seed,
                parameterization=parameterization,
            )
        if name == "f2_square":
            return F2SquareAttack(
                epsilon_mv=epsilon_mv,
                initial_fraction=float(values.get("initial_fraction", 0.20)),
                norm=norm,
                seed=seed,
                parameterization=parameterization,
                parameter_amplitude=float(values.get("parameter_amplitude", 1.0)),
            )
        knowledge = {
            "f3_k1": KnowledgeLevel.K1,
            "f3_k2": KnowledgeLevel.K2,
            "f3_k3": KnowledgeLevel.K3,
        }[name]
        return F3AdaptiveAttack(
            knowledge_level=knowledge,
            diagnosis_weight=float(values.get("diagnosis_weight", 1.0)),
            retrieval_weight=float(values.get("retrieval_weight", 1.0)),
            verifier_weight=float(values.get("verifier_weight", 1.0)),
            defense_weight=float(values.get("defense_weight", 0.0)),
            black_box_method=str(values.get("black_box_method", "nes")),
            black_box_parameterization=parameterization,
            **gradient,
        )

    family = {
        "f0": AttackFamily.F0,
        "f1_pgd": AttackFamily.F1,
        "f1_cf": AttackFamily.F1_CF,
        "f2_nes": AttackFamily.F2,
        "f2_simba": AttackFamily.F2,
        "f2_square": AttackFamily.F2,
        "f3_k1": AttackFamily.F3,
        "f3_k2": AttackFamily.F3,
        "f3_k3": AttackFamily.F3,
    }[name]
    return family, factory


def _security_admissibility(values: Mapping[str, Any], defense: str):
    from .security.measurements import (
        ClinicalAdmissibilityChecker,
        NumpyMeasurementBackend,
        ScipyMeasurementBackend,
    )

    if not bool(values.get("enabled", True)):
        return None
    backend_name = str(values.get("backend", "scipy"))
    if backend_name == "scipy":
        backend = ScipyMeasurementBackend()
    elif backend_name == "numpy":
        if defense == "v2":
            raise ValueError(
                "V2 and admissibility cannot both use the NumPy measurement implementation"
            )
        backend = NumpyMeasurementBackend()
    else:
        raise ValueError("Admissibility backend must be scipy or numpy")
    kwargs: dict[str, Any] = {"backend": backend}
    if "tolerances" in values:
        kwargs["tolerances"] = {
            str(key): float(value) for key, value in values["tolerances"].items()
        }
    if "required" in values:
        kwargs["required"] = tuple(str(value) for value in values["required"])
    return ClinicalAdmissibilityChecker(**kwargs)


def _security_victim_provider(
    pipeline: CAREECGPipeline,
    *,
    query: str,
    defense: str,
):
    from .security.contracts import DefenseKind
    from .security.defenses import MeasurementAwareVerifier
    from .security.measurements import NumpyMeasurementBackend
    from .security.victim import InternalCAREECGVictim

    defense_kind = DefenseKind(defense)
    evaluators: dict[DefenseKind, Any] = {}
    if defense_kind is DefenseKind.V2:
        evaluators[DefenseKind.V2] = MeasurementAwareVerifier(backend=NumpyMeasurementBackend())

    def provider(*, case: Any, **_: Any):
        return InternalCAREECGVictim(
            pipeline,
            query=query,
            patient_id=case.patient_id,
            defense_evaluators=evaluators,
        )

    return provider


def _security_lock_payload(
    security_config: str | Path,
    validation_results: str | Path,
) -> dict[str, Any]:
    config_path = Path(security_config).resolve()
    results_path = Path(validation_results).resolve()
    report = json.loads(results_path.read_text(encoding="utf-8"))
    if report.get("schema_version") != SECURITY_RESULTS_SCHEMA:
        raise ValueError("Validation results use an unsupported security schema")
    report_protocol = report.get("protocol", {})
    if report_protocol.get("phase") != "validation":
        raise ValueError("A security lock can only be created from a validation-phase result")
    config_sha256 = _sha256_file(config_path)
    if report_protocol.get("security_config_sha256") != config_sha256:
        raise ValueError("Validation results were produced with a different security config")
    security_config = _load_security_config(config_path)
    expected_folds = sorted(
        int(value) for value in security_config["protocol"].get("validation_folds", (9,))
    )
    if report_protocol.get("observed_folds") != expected_folds:
        raise ValueError("Validation results do not contain the predeclared validation folds")
    candidate_cases = report_protocol.get("candidate_cases")
    eligible_cases = report_protocol.get("eligible_cases_used")
    if (
        isinstance(candidate_cases, bool)
        or not isinstance(candidate_cases, int)
        or candidate_cases < 1
        or isinstance(eligible_cases, bool)
        or not isinstance(eligible_cases, int)
        or eligible_cases < 1
    ):
        raise ValueError("Validation lock requires a nonempty completed attack cohort")

    e0 = report.get("e0")
    e1 = report.get("e1")
    if not isinstance(e0, Mapping) or len(e0.get("observations", ())) != candidate_cases:
        raise ValueError("Validation results contain an incomplete E0 cohort")
    if not isinstance(e1, Mapping):
        raise TypeError("Validation results do not contain an E1 matrix")
    for attack_name in report_protocol.get("attacks", ()):
        by_constraint = e1.get(attack_name)
        if not isinstance(by_constraint, Mapping):
            raise TypeError(f"Validation results omit E1 attack {attack_name}")
        for constraint_name in report_protocol.get("constraints", ()):
            run = by_constraint.get(constraint_name)
            if not isinstance(run, Mapping) or not run.get("observations"):
                raise ValueError(
                    f"Validation results omit E1 observations for {attack_name}/{constraint_name}"
                )

    event_log = report.get("event_log")
    if not isinstance(event_log, Mapping):
        raise TypeError("Validation results do not contain an event-log manifest")
    log_path = Path(str(event_log.get("path", "")))
    manifest_path = Path(str(event_log.get("manifest_path", "")))
    if not log_path.is_file() or _sha256_file(log_path) != event_log.get("sha256"):
        raise ValueError("Validation event log is unavailable or changed")
    if not manifest_path.is_file() or _sha256_file(manifest_path) != event_log.get(
        "manifest_sha256"
    ):
        raise ValueError("Validation terminal manifest is unavailable or changed")
    from .security.logging import AppendOnlySecurityLogger

    verified_log = AppendOnlySecurityLogger(
        log_path,
        manifest_path=manifest_path,
    ).verify(require_manifest=True)
    if not verified_log.valid or verified_log.complete is not True:
        raise ValueError(
            "Validation event log is invalid or incomplete: " + "; ".join(verified_log.errors)
        )
    recorded_verification = event_log.get("verification")
    if not isinstance(recorded_verification, Mapping) or not (
        recorded_verification.get("valid") is True
        and recorded_verification.get("complete") is True
        and recorded_verification.get("records") == verified_log.records
        and recorded_verification.get("terminal_record_sha256")
        == verified_log.terminal_record_sha256
    ):
        raise ValueError("Validation report does not match its independently verified event log")

    artifact_identity = report.get("artifacts", {}).get("identity")
    code_identity = report.get("implementation", {}).get("code_identity")
    preprocessing = report.get("dataset", {}).get("preprocessing")
    if not isinstance(artifact_identity, dict) or not artifact_identity.get("sha256"):
        raise ValueError("Validation results do not bind the victim artifact identity")
    if not isinstance(code_identity, dict) or not code_identity.get("sha256"):
        raise ValueError("Validation results do not bind the security code identity")
    if not isinstance(preprocessing, dict):
        raise TypeError("Validation results do not bind preprocessing settings")
    artifact_path = Path(str(report.get("artifacts", {}).get("path", "")))
    if _directory_identity(artifact_path) != artifact_identity:
        raise ValueError("Validation victim artifact tree is unavailable or changed")
    if _security_code_identity() != code_identity:
        raise ValueError("Security implementation changed before validation was locked")
    dataset = report.get("dataset", {})
    dataset_path = Path(str(dataset.get("path", "")))
    if not dataset_path.is_file() or _sha256_file(dataset_path) != dataset.get("sha256"):
        raise ValueError("Validation dataset is unavailable or changed")
    expected_dataset_hash = security_config["protocol"].get("validation_dataset_sha256")
    if expected_dataset_hash != dataset.get("sha256"):
        raise ValueError("Validation dataset does not match the predeclared protocol hash")
    frozen_fields = (
        "target_ids",
        "budgets",
        "seeds",
        "attacks",
        "constraints",
        "defense",
        "max_candidate_patients_per_target",
        "max_eligible_patients_per_target",
        "record_level_exploratory",
        "allow_missing_folds",
        "source_rate_hz",
        "e0_report_correctness_endpoint",
        "clinical_preservation_endpoint",
        "f0_magnitude_calibration",
    )
    resolved_protocol = {
        key: report["protocol"].get(key) for key in frozen_fields if key in report["protocol"]
    }
    if "f0" in resolved_protocol.get("attacks", ()):
        calibration = resolved_protocol.get("f0_magnitude_calibration")
        if not isinstance(calibration, Mapping):
            raise TypeError("Validation F0 results do not include magnitude calibration")
        for constraint_name in resolved_protocol.get("constraints", ()):
            _f0_magnitudes_for_constraint(
                calibration,
                str(constraint_name),
                tuple(int(value) for value in resolved_protocol.get("budgets", ())),
            )
    return {
        "schema_version": SECURITY_LOCK_SCHEMA,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "security_config_path": str(config_path),
        "security_config_sha256": config_sha256,
        "validation_results_path": str(results_path),
        "validation_results_sha256": _sha256_file(results_path),
        "resolved_protocol": resolved_protocol,
        "artifact_identity": artifact_identity,
        "code_identity": code_identity,
        "preprocessing": preprocessing,
    }


def _security_lock(args: argparse.Namespace) -> int:
    payload = _security_lock_payload(args.security_config, args.validation_results)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"validation_lock": str(output.resolve()), **payload}, indent=2))
    return 0


def _validate_security_lock(path: str | Path, security_config: str | Path) -> dict[str, Any]:
    lock_path = Path(path)
    payload = json.loads(lock_path.read_text(encoding="utf-8"))
    if payload.get("schema_version") != SECURITY_LOCK_SCHEMA:
        raise ValueError("Held-out evaluation requires a supported validation lock")
    if payload.get("security_config_sha256") != _sha256_file(security_config):
        raise ValueError("Security config changed after validation was locked")
    validation_path = Path(str(payload.get("validation_results_path", "")))
    if not validation_path.is_file():
        raise ValueError("Locked validation results are unavailable")
    if payload.get("validation_results_sha256") != _sha256_file(validation_path):
        raise ValueError("Locked validation results changed after the lock was created")
    for field in ("artifact_identity", "code_identity", "preprocessing"):
        if not isinstance(payload.get(field), dict):
            raise TypeError(f"Validation lock does not bind {field.replace('_', ' ')}")
    return {
        **payload,
        "lock_path": str(lock_path.resolve()),
        "lock_sha256": _sha256_file(lock_path),
    }


def _excel_cell(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return json.dumps(value, sort_keys=True, separators=(",", ":"))


def _security_paired_f3_f1(e1: Mapping[str, Mapping[str, Any]]) -> dict[str, Any]:
    from .security.metrics import bootstrap_paired_irr_difference

    if "f1_pgd" not in e1:
        return {}
    comparison: dict[str, Any] = {}
    for adaptive_name in ("f3_k1", "f3_k2", "f3_k3"):
        if adaptive_name not in e1:
            continue
        shared_constraints = sorted(set(e1["f1_pgd"]) & set(e1[adaptive_name]))
        for constraint_name in shared_constraints:
            f1_rows = e1["f1_pgd"][constraint_name]["observations"]
            adaptive_rows = e1[adaptive_name][constraint_name]["observations"]

            def index_rows(
                rows: Sequence[Mapping[str, Any]], arm: str, constraint_label: str
            ) -> dict[tuple[Any, ...], Mapping[str, Any]]:
                indexed: dict[tuple[Any, ...], Mapping[str, Any]] = {}
                for row in rows:
                    key = (
                        row["patient_id"],
                        row["record_id"],
                        row["target_id"],
                        row["seed"],
                        row["budget"],
                        row["defense"],
                    )
                    if key in indexed:
                        raise ValueError(
                            f"Duplicate paired observation for {arm}/{constraint_label}: {key}"
                        )
                    indexed[key] = row
                return indexed

            f1_index = index_rows(f1_rows, "f1_pgd", constraint_name)
            adaptive_index = index_rows(adaptive_rows, adaptive_name, constraint_name)
            if f1_index.keys() != adaptive_index.keys():
                missing_adaptive = len(f1_index.keys() - adaptive_index.keys())
                missing_comparator = len(adaptive_index.keys() - f1_index.keys())
                raise ValueError(
                    f"Incomplete {adaptive_name}/f1_pgd pairing for {constraint_name}: "
                    f"{missing_adaptive} missing adaptive and "
                    f"{missing_comparator} missing comparator observations"
                )

            grouped: dict[
                tuple[str, int, str],
                list[tuple[Mapping[str, Any], Mapping[str, Any]]],
            ] = {}
            for identity, row in adaptive_index.items():
                group = (
                    str(row["target_id"]),
                    int(row["budget"]),
                    str(row["defense"]),
                )
                grouped.setdefault(group, []).append((row, f1_index[identity]))
            for (target_id, budget, defense), pairs in grouped.items():
                primary_pairs = [
                    (adaptive, comparator)
                    for adaptive, comparator in pairs
                    if bool((adaptive.get("metrics") or {}).get("primary_endpoint_eligible", False))
                    and bool(
                        (comparator.get("metrics") or {}).get("primary_endpoint_eligible", False)
                    )
                ]
                metadata = {
                    "target_id": target_id,
                    "paired_observations": len(primary_pairs),
                    "total_paired_observations": len(pairs),
                    "excluded_nonprimary_pairs": len(pairs) - len(primary_pairs),
                    "seeds": sorted({int(adaptive["seed"]) for adaptive, _ in pairs}),
                    "endpoint_classification": "primary",
                }
                if primary_pairs:
                    estimate = bootstrap_paired_irr_difference(
                        [bool(adaptive["success"]) for adaptive, _ in primary_pairs],
                        [bool(comparator["success"]) for _, comparator in primary_pairs],
                        [str(adaptive["patient_id"]) for adaptive, _ in primary_pairs],
                        n_resamples=500,
                        seed=42,
                    ).to_dict()
                    estimate["metadata"].update({**metadata, "status": "evaluable"})
                else:
                    estimate = {
                        "estimate": None,
                        "lower": None,
                        "upper": None,
                        "confidence": 0.95,
                        "n_patients": 0,
                        "n_resamples": 0,
                        "seed": 42,
                        "metadata": {**metadata, "status": "not_evaluable"},
                    }
                key = f"{adaptive_name}-f1_pgd:{constraint_name}|T={target_id}|B={budget}|{defense}"
                comparison[key] = estimate
    return comparison


def _f0_magnitude_calibration(
    e1: Mapping[str, Mapping[str, Any]],
    *,
    constraints: Sequence[str],
    budgets: Sequence[int],
    source_attacks: Sequence[str],
) -> dict[str, Any]:
    """Freeze median native-set distances from validation non-F0 submissions."""

    selected_sources = tuple(attack for attack in source_attacks if attack != "f0" and attack in e1)
    if not selected_sources:
        raise ValueError(
            "F0 validation calibration requires at least one selected non-F0 source attack"
        )
    distance_units = {
        "s_inf": "max_abs_mv",
        "s2": "max_per_lead_l2_mv",
        "s_spec": "max_abs_mv_with_spectral_feasibility",
        "s_phys": "max_normalized_parameter_displacement",
    }
    strata: dict[str, dict[str, Any]] = {}
    for constraint_name in constraints:
        for budget in budgets:
            distances: list[float] = []
            contributing: set[str] = set()
            for attack_name in selected_sources:
                run = e1.get(attack_name, {}).get(constraint_name)
                if run is None:
                    continue
                for observation in run.get("observations", ()):  # type: ignore[union-attr]
                    if int(observation.get("budget", -1)) != int(budget):
                        continue
                    attack_result = observation.get("attack_result") or {}
                    compact_distances = attack_result.get("constraint_distances", ())
                    compact_satisfied = attack_result.get("constraint_satisfied", ())
                    if len(compact_distances) != len(compact_satisfied):
                        raise ValueError("F0 calibration trace lengths do not match")
                    for value, satisfied in zip(compact_distances, compact_satisfied, strict=True):
                        if not bool(satisfied):
                            continue
                        if value is None:
                            raise ValueError(
                                "Attack attempt omits native constraint_distance required "
                                "for F0 calibration"
                            )
                        numeric = float(value)
                        if not np.isfinite(numeric) or numeric < 0:
                            raise ValueError(
                                "F0 calibration distance must be finite and non-negative"
                            )
                        distances.append(numeric)
                        contributing.add(attack_name)
            if not distances:
                raise ValueError(
                    f"F0 calibration stratum {constraint_name}/B={budget} has no feasible "
                    "non-F0 validation submissions"
                )
            key = f"{constraint_name}|B={int(budget)}"
            strata[key] = {
                "constraint": constraint_name,
                "budget": int(budget),
                "distance_unit": distance_units[constraint_name],
                "statistic": "median",
                "magnitude": float(np.median(np.asarray(distances, dtype=float))),
                "sample_count": len(distances),
                "source_attacks": sorted(contributing),
                "minimum": float(np.min(distances)),
                "maximum": float(np.max(distances)),
            }
    calibration: dict[str, Any] = {
        "schema_version": "care-ecg-f0-magnitude-calibration.v1",
        "phase": "validation",
        "reference_pool": "all feasible submitted queries from selected non-F0 arms",
        "source_attacks": list(selected_sources),
        "stratification": ["constraint", "budget"],
        "statistic": "median",
        "strata": strata,
    }
    calibration["calibration_sha256"] = hashlib.sha256(
        json.dumps(calibration, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    return calibration


def _f0_magnitudes_for_constraint(
    calibration: Mapping[str, Any], constraint_name: str, budgets: Sequence[int]
) -> dict[int, float]:
    if calibration.get("schema_version") != "care-ecg-f0-magnitude-calibration.v1":
        raise ValueError("F0 magnitude calibration uses an unsupported schema")
    unsigned = dict(calibration)
    supplied_hash = unsigned.pop("calibration_sha256", None)
    observed_hash = hashlib.sha256(
        json.dumps(unsigned, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    if supplied_hash != observed_hash:
        raise ValueError("F0 magnitude calibration hash mismatch")
    strata = calibration.get("strata")
    if not isinstance(strata, Mapping):
        raise TypeError("F0 magnitude calibration has no strata")
    result: dict[int, float] = {}
    for budget in budgets:
        item = strata.get(f"{constraint_name}|B={int(budget)}")
        if not isinstance(item, Mapping):
            raise TypeError(f"F0 magnitude calibration omits {constraint_name}/B={int(budget)}")
        magnitude = float(item.get("magnitude", float("nan")))
        if not np.isfinite(magnitude) or magnitude < 0:
            raise ValueError("F0 matched magnitude must be finite and non-negative")
        result[int(budget)] = magnitude
    return result


def _write_security_workbook(report: Mapping[str, Any], output: str | Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise RuntimeError("Excel export requires pip install -e '.[reports]'") from exc

    workbook = Workbook()
    manifest_sheet = workbook.active
    manifest_sheet.title = "Manifest"
    manifest_sheet.append(("field", "value"))
    for key, value in report.items():
        if key not in {"e0", "e1", "paired_f3_minus_f1"}:
            manifest_sheet.append((key, _excel_cell(value)))

    summary_sheet = workbook.create_sheet("Summaries")
    summary_sheet.append(("experiment", "run", "summary"))
    summary_sheet.append(("E0", "clean", _excel_cell(report["e0"]["summary"])))
    for attack_name, by_constraint in report["e1"].items():
        for constraint_name, run in by_constraint.items():
            summary_sheet.append(
                ("E1", f"{attack_name}:{constraint_name}", _excel_cell(run["summary"]))
            )
    summary_sheet.append(("E1", "paired_f3_minus_f1", _excel_cell(report["paired_f3_minus_f1"])))

    robustness_sheet = workbook.create_sheet("Robustness metrics")
    robustness_sheet.append(
        (
            "attack",
            "constraint",
            "stratum",
            "metric",
            "estimate",
            "ci_lower",
            "ci_upper",
            "confidence",
            "patients",
            "observations",
            "clean_correct_observations",
            "endpoint_classification",
        )
    )
    distortion_sheet = workbook.create_sheet("Distortion metrics")
    distortion_sheet.append(
        (
            "attack",
            "constraint",
            "stratum",
            "cohort",
            "metric",
            "count",
            "mean",
            "median",
            "p95",
            "max",
        )
    )
    descriptive_sheet = workbook.create_sheet("Attack descriptives")
    descriptive_sheet.append(
        (
            "attack",
            "constraint",
            "stratum",
            "metric",
            "count",
            "mean",
            "median",
            "p95",
            "max",
        )
    )
    estimate_names = (
        "clean_accuracy",
        "robust_accuracy",
        "clean_to_robust_accuracy_drop",
        "targeted_diagnosis_attack_success_rate",
        "untargeted_diagnosis_attack_success_rate",
        "released_target_claim_attack_success_rate",
        "valid_selected_perturbation_rate",
    )
    descriptive_names = (
        "queries_to_first_targeted_diagnosis",
        "queries_to_first_released_target_claim",
        "target_probability_gain",
        "target_log_probability_gain",
    )
    distortion_groups = (
        "distortion_all_valid_selected",
        "distortion_targeted_diagnosis_successes",
    )
    for attack_name, by_constraint in report["e1"].items():
        for constraint_name, run in by_constraint.items():
            for stratum, stratum_summary in run["summary"].items():
                if not isinstance(stratum_summary, Mapping):
                    continue
                automated = stratum_summary.get("automated_robustness")
                if not isinstance(automated, Mapping):
                    continue
                base = (
                    attack_name,
                    constraint_name,
                    stratum,
                )
                for metric_name in estimate_names:
                    estimate = automated.get(metric_name)
                    if not isinstance(estimate, Mapping):
                        continue
                    robustness_sheet.append(
                        (
                            *base,
                            metric_name,
                            estimate.get("estimate"),
                            estimate.get("lower"),
                            estimate.get("upper"),
                            estimate.get("confidence"),
                            estimate.get("n_patients"),
                            automated.get("observations"),
                            automated.get("clean_correct_observations"),
                            automated.get("endpoint_classification"),
                        )
                    )
                for budget_group in (
                    "targeted_diagnosis_asr_at_budget",
                    "robust_accuracy_at_budget",
                ):
                    estimates = automated.get(budget_group)
                    if not isinstance(estimates, Mapping):
                        continue
                    for budget, estimate in estimates.items():
                        if not isinstance(estimate, Mapping):
                            continue
                        robustness_sheet.append(
                            (
                                *base,
                                f"{budget_group}@{budget}",
                                estimate.get("estimate"),
                                estimate.get("lower"),
                                estimate.get("upper"),
                                estimate.get("confidence"),
                                estimate.get("n_patients"),
                                automated.get("observations"),
                                automated.get("clean_correct_observations"),
                                automated.get("endpoint_classification"),
                            )
                        )
                for metric_name in descriptive_names:
                    values = automated.get(metric_name)
                    if not isinstance(values, Mapping):
                        continue
                    descriptive_sheet.append(
                        (
                            *base,
                            metric_name,
                            values.get("count"),
                            values.get("mean"),
                            values.get("median"),
                            values.get("p95"),
                            values.get("max"),
                        )
                    )
                for cohort in distortion_groups:
                    distortions = automated.get(cohort)
                    if not isinstance(distortions, Mapping):
                        continue
                    for metric_name, values in distortions.items():
                        if not isinstance(values, Mapping):
                            continue
                        distortion_sheet.append(
                            (
                                *base,
                                cohort,
                                metric_name,
                                values.get("count"),
                                values.get("mean"),
                                values.get("median"),
                                values.get("p95"),
                                values.get("max"),
                            )
                        )

    e0_sheet = workbook.create_sheet("E0 observations")
    observation_fields = (
        "experiment_id",
        "arm",
        "patient_id",
        "record_id",
        "target_id",
        "success",
        "released",
        "defense",
        "metrics",
        "metadata",
    )
    e0_sheet.append(observation_fields)
    for observation in report["e0"]["observations"]:
        e0_sheet.append(tuple(_excel_cell(observation.get(key)) for key in observation_fields))

    e1_sheet = workbook.create_sheet("E1 observations")
    e1_fields = (
        "attack",
        "constraint",
        "patient_id",
        "record_id",
        "target_id",
        "success",
        "released",
        "seed",
        "budget",
        "defense",
        "queries_used",
        "metrics",
        "metadata",
        "attack_result",
    )
    e1_sheet.append(e1_fields)
    for attack_name, by_constraint in report["e1"].items():
        for constraint_name, run in by_constraint.items():
            for observation in run["observations"]:
                values = {**observation, "attack": attack_name, "constraint": constraint_name}
                e1_sheet.append(tuple(_excel_cell(values.get(key)) for key in e1_fields))

    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    destination = Path(output)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)


def _security_evaluate(args: argparse.Namespace) -> int:
    from .security.contracts import DefenseKind
    from .security.experiments import run_e0_clean_baseline, run_e1_attack_matrix
    from .security.logging import AppendOnlySecurityLogger

    security_config_path = Path(args.security_config).resolve()
    security_config = _load_security_config(security_config_path)
    protocol = security_config["protocol"]
    excel_enabled = (
        bool(security_config["export"].get("excel", True))
        if args.excel is None
        else bool(args.excel)
    )
    if excel_enabled:
        try:
            __import__("openpyxl")
        except ImportError as exc:  # pragma: no cover - optional dependency
            raise RuntimeError("Excel export requires pip install -e '.[reports]'") from exc
    phase = args.phase or str(protocol.get("phase", "validation"))
    if phase not in {"validation", "heldout"}:
        raise ValueError("Security phase must be validation or heldout")
    if phase == "heldout" and args.allow_missing_folds:
        raise ValueError("Held-out primary evaluation cannot bypass fold metadata")
    validation_lock = None
    if phase == "heldout":
        if not args.validation_lock:
            raise ValueError("Held-out evaluation requires --validation-lock")
        validation_lock = _validate_security_lock(args.validation_lock, security_config_path)

    artifact_identity = _directory_identity(args.artifacts)
    code_identity = _security_code_identity()
    if validation_lock is not None:
        if validation_lock["artifact_identity"] != artifact_identity:
            raise ValueError("Victim artifacts differ from the validation-locked identity")
        if validation_lock["code_identity"] != code_identity:
            raise ValueError("CARE-ECG code differs from the validation-locked identity")

    pipeline = CAREECGPipeline.load(args.artifacts)
    dataset_path = Path(args.dataset).resolve()
    dataset_sha256 = _sha256_file(dataset_path)
    expected_dataset_hash = protocol.get(f"{phase}_dataset_sha256")
    if expected_dataset_hash is None:
        raise ValueError(f"Security protocol must predeclare protocol.{phase}_dataset_sha256")
    if str(expected_dataset_hash) != dataset_sha256:
        raise ValueError(f"{phase.capitalize()} dataset hash differs from the protocol")
    bundle = load_npz(dataset_path, require_labels=True)
    allowed_folds = tuple(
        int(value)
        for value in protocol.get(
            "validation_folds" if phase == "validation" else "heldout_folds",
            (9,) if phase == "validation" else (10,),
        )
    )
    observed_folds: list[int] | None = None
    if bundle.folds is None:
        if not args.allow_missing_folds:
            raise ValueError("Security evaluation requires fold metadata")
    else:
        observed_folds = sorted({int(value) for value in bundle.folds.tolist()})
        unexpected = sorted(set(observed_folds) - set(allowed_folds))
        if unexpected:
            raise ValueError(
                f"{phase} security run contains disallowed folds {unexpected}; "
                f"expected only {list(allowed_folds)}"
            )

    attacks = _csv_items(
        args.attacks,
        security_config["attacks"].get("enabled", SECURITY_ATTACK_NAMES),
    )
    constraints = _csv_items(
        args.constraints,
        security_config["constraints"].get("enabled", SECURITY_CONSTRAINT_NAMES),
    )
    target_ids = _csv_items(args.target_ids, protocol.get("targets", ("T1", "T2", "T3", "T4")))
    budgets = _csv_items(args.budgets, protocol.get("budgets", (1, 10, 50, 100)), int)
    seeds = _csv_items(args.seeds, protocol.get("seeds", (42,)), int)
    unknown_attacks = sorted(set(attacks) - set(SECURITY_ATTACK_NAMES))
    unknown_constraints = sorted(set(constraints) - set(SECURITY_CONSTRAINT_NAMES))
    if unknown_attacks:
        raise ValueError(f"Unknown security attacks: {unknown_attacks}")
    if unknown_constraints:
        raise ValueError(f"Unknown security constraints: {unknown_constraints}")
    if any(value < 1 for value in budgets):
        raise ValueError("Security query budgets must be positive")
    if (
        "f0" in attacks
        and security_config["attacks"].get("f0", {}).get("matched_magnitude") is not None
    ):
        raise ValueError(
            "Protocol F0 magnitude must be validation-derived; remove the fixed "
            "attacks.f0.matched_magnitude override"
        )

    signals = _security_raw_signals(bundle.signals, pipeline, args.source_rate)
    preprocessing = {
        "input_units": "mV",
        "source_rate_hz": int(args.source_rate or pipeline.config.data.sample_rate),
        "target_rate_hz": int(pipeline.config.data.sample_rate),
        "seconds": int(pipeline.config.data.seconds),
        "leads": int(pipeline.config.data.leads),
        "normalization": "none_before_victim_artifact_preprocessing",
    }
    if validation_lock is not None and validation_lock["preprocessing"] != preprocessing:
        raise ValueError("Preprocessing differs from the validation-locked contract")
    candidate_limit = args.max_candidates_per_target
    if candidate_limit is None:
        configured_limit = protocol.get("max_candidate_patients_per_target")
        candidate_limit = None if configured_limit is None else int(configured_limit)
    cases = _security_cohort_cases(
        bundle,
        signals,
        pipeline.causal_model.label_names,
        target_ids,
        max_candidates_per_target=candidate_limit,
        allow_record_level=args.allow_record_level,
    )
    defense = args.defense or str(protocol.get("defense", "v1"))
    if defense not in {DefenseKind.V1.value, DefenseKind.V2.value}:
        raise ValueError("Artifact-path security evaluation supports defense v1 or v2")
    victim_provider = _security_victim_provider(
        pipeline,
        query=str(protocol.get("query", "Which diagnosis is most supported, and why?")),
        defense=defense,
    )
    admissibility = _security_admissibility(security_config["admissibility"], defense)
    eligible_limit = int(
        args.max_eligible_per_target
        if args.max_eligible_per_target is not None
        else protocol.get("max_eligible_patients_per_target", 100)
    )
    resolved_protocol = {
        "target_ids": list(target_ids),
        "budgets": list(budgets),
        "seeds": list(seeds),
        "attacks": list(attacks),
        "constraints": list(constraints),
        "defense": defense,
        "max_candidate_patients_per_target": candidate_limit,
        "max_eligible_patients_per_target": eligible_limit,
        "record_level_exploratory": bool(args.allow_record_level),
        "allow_missing_folds": bool(args.allow_missing_folds),
        "source_rate_hz": preprocessing["source_rate_hz"],
        "e0_report_correctness_endpoint": ("exploratory_diagnosis_and_target_absence_heuristic"),
        "clinical_preservation_endpoint": "exploratory_local_delta_only",
    }
    f0_calibration: Mapping[str, Any] | None = None
    if "f0" in attacks and phase == "heldout":
        locked_protocol = validation_lock.get("resolved_protocol", {})
        f0_calibration = locked_protocol.get("f0_magnitude_calibration")
        if not isinstance(f0_calibration, Mapping):
            raise TypeError("Held-out F0 evaluation requires locked validation calibration")
        for constraint_name in constraints:
            _f0_magnitudes_for_constraint(f0_calibration, constraint_name, budgets)
        resolved_protocol["f0_magnitude_calibration"] = dict(f0_calibration)
    if (
        validation_lock is not None
        and validation_lock.get("resolved_protocol")
        and validation_lock["resolved_protocol"] != resolved_protocol
    ):
        raise ValueError(
            "Held-out protocol differs from the validation-locked targets, attacks, "
            "constraints, budgets, seeds, defense, or cohort limits"
        )

    output_dir = Path(args.output)
    output_dir.mkdir(parents=True, exist_ok=True)
    result_path = output_dir / "security_results.json"
    log_path = output_dir / "security_events.jsonl"
    manifest_path = Path(f"{log_path}.manifest.json")
    if result_path.exists() or log_path.exists() or manifest_path.exists():
        raise FileExistsError(
            "Security output already contains results, an event log, or a terminal manifest; "
            "use a new output directory"
        )
    logger = AppendOnlySecurityLogger(log_path)
    print(f"Running E0 on {len(cases)} patient-target cases", file=sys.stderr)
    e0 = run_e0_clean_baseline(
        cases,
        victim_provider,
        defense=defense,
        logger=logger,
    )
    eligible = _limit_eligible_cases(e0.eligible_cases, eligible_limit)

    e1: dict[str, dict[str, Any]] = {}
    constraint_section = security_config["constraints"]
    execution_order = tuple(name for name in attacks if name != "f0") + (
        ("f0",) if "f0" in attacks else ()
    )
    for attack_name in execution_order:
        if attack_name == "f0" and f0_calibration is None:
            f0_values = security_config["attacks"].get("f0", {})
            if not isinstance(f0_values, Mapping):
                raise TypeError("attacks.f0 must be a mapping")
            if str(f0_values.get("match_statistic", "median")) != "median":
                raise ValueError("The supported F0 validation match statistic is median")
            source_attacks = tuple(
                str(value)
                for value in f0_values.get(
                    "match_reference_attacks",
                    tuple(name for name in attacks if name != "f0"),
                )
            )
            f0_calibration = _f0_magnitude_calibration(
                e1,
                constraints=constraints,
                budgets=budgets,
                source_attacks=source_attacks,
            )
            resolved_protocol["f0_magnitude_calibration"] = dict(f0_calibration)
        e1[attack_name] = {}
        for constraint_name in constraints:
            constraint_values = constraint_section.get(constraint_name, {})
            if not isinstance(constraint_values, dict):
                raise TypeError(f"Constraint {constraint_name!r} configuration must be a mapping")
            constraint = _security_constraint(
                constraint_name,
                constraint_values,
                float(pipeline.config.data.sample_rate),
            )
            family, factory = _security_attack_factory(
                attack_name,
                security_config["attacks"],
                constraint,
                matched_magnitudes=(
                    None
                    if attack_name != "f0" or f0_calibration is None
                    else _f0_magnitudes_for_constraint(f0_calibration, constraint_name, budgets)
                ),
            )
            print(
                f"Running E1 {attack_name}/{constraint_name} on {len(eligible)} eligible cases",
                file=sys.stderr,
            )
            run = run_e1_attack_matrix(
                eligible,
                victim_provider,
                {family: factory},
                {constraint_name: constraint},
                budgets=budgets,
                seeds=seeds,
                defense=defense,
                admissibility_checker=admissibility,
                logger=logger,
            )
            e1[attack_name][constraint_name] = run.to_dict()

    terminal_manifest = logger.seal()
    log_verification = logger.verify(require_manifest=True).to_dict()
    if not log_verification["valid"] or log_verification["complete"] is not True:
        raise RuntimeError(
            "Security event log did not pass terminal-manifest verification: "
            + "; ".join(log_verification["errors"])
        )
    artifact_path = Path(args.artifacts).resolve()
    dataset_path = Path(args.dataset).resolve()
    config_hash = _sha256_file(security_config_path)
    report: dict[str, Any] = {
        "schema_version": SECURITY_RESULTS_SCHEMA,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "protocol": {
            "phase": phase,
            "security_config_path": str(security_config_path),
            "security_config_sha256": config_hash,
            "validation_lock": validation_lock,
            "allowed_folds": list(allowed_folds),
            "observed_folds": observed_folds,
            **resolved_protocol,
            "candidate_cases": len(cases),
            "eligible_cases_used": len(eligible),
        },
        "artifacts": {
            "path": str(artifact_path),
            "config_sha256": _sha256_file(artifact_path / "config.yaml"),
            "model_sha256": _sha256_file(artifact_path / "causal_vae.pt"),
            "identity": artifact_identity,
            "victim_scope": "waveform-callable internal CARE model",
            "external_record_keyed_expert_attacked": False,
        },
        "dataset": {
            "path": str(dataset_path),
            "sha256": dataset_sha256,
            "protocol_sha256": str(expected_dataset_hash),
            "input_units": "mV (asserted by the security protocol)",
            "patient_ids_available": bundle.patient_ids is not None,
            "preprocessing": preprocessing,
            "candidate_cohort": _security_cohort_manifest(cases),
            "eligible_cohort": _security_cohort_manifest([item.case for item in eligible]),
        },
        "implementation": {
            "package_version": __version__,
            "code_identity": code_identity,
        },
        "admissibility": {
            "enabled": admissibility is not None,
            "scope": (
                "exploratory local delta implementation; no target-specific "
                "preserved-finding reference evidence"
            ),
            "backend": None if admissibility is None else admissibility.backend.name,
            "endpoint_classification": "exploratory",
            "primary_endpoint_eligible": False,
        },
        "event_log": {
            "path": str(log_path.resolve()),
            "sha256": _sha256_file(log_path),
            "manifest_path": str(logger.manifest_path.resolve()),
            "manifest_sha256": _sha256_file(logger.manifest_path),
            "terminal_manifest": terminal_manifest,
            "verification": log_verification,
        },
        "e0": e0.run.to_dict(),
        "e1": e1,
        "paired_f3_minus_f1": _security_paired_f3_f1(e1),
        "limitations": [
            "The cached record-keyed external ensemble cannot be recomputed after waveform edits and is excluded from the attacked victim.",
            "The bundled local clinical-delta checker and E0 report heuristic are exploratory and cannot contribute primary-endpoint successes.",
            "Primary evaluation requires independent all-claims clean-report review and target-specific preserved-finding evidence through the Python API.",
            "Adult-only eligibility must be applied upstream because the standard NPZ contract has no age field.",
            "This command runs E0/E1; E1b-E7 require their explicit intervention, transfer-victim, calibration, or authenticated-history Python callbacks.",
        ],
    }
    result_path.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")

    workbook_path = output_dir / "security_results.xlsx"
    if excel_enabled:
        _write_security_workbook(report, workbook_path)
    result = {
        "results_json": str(result_path.resolve()),
        "results_xlsx": str(workbook_path.resolve()) if excel_enabled else None,
        "event_log": str(log_path.resolve()),
        "event_log_valid": log_verification["valid"],
        "event_log_complete": log_verification["complete"],
        "event_log_manifest": str(logger.manifest_path.resolve()),
        "analysis_eligible_cases": len(eligible),
        "primary_eligible_cases": sum(item.primary_endpoint_eligible for item in eligible),
        "endpoint_classification": "exploratory",
    }
    print(json.dumps(result, indent=2))
    return 0


def _security_verify_log(args: argparse.Namespace) -> int:
    from .security.logging import AppendOnlySecurityLogger

    logger = AppendOnlySecurityLogger(args.log, args.waveform_dir)
    result = logger.verify(require_manifest=True).to_dict()
    print(json.dumps(result, indent=2))
    return 0 if result["valid"] else 1


def _demo(args: argparse.Namespace) -> int:
    config = CAREConfig()
    config.seed = args.seed
    config.data.sample_rate = 100
    config.data.seconds = 2
    config.model.latent_dim = 4
    config.model.hidden_dim = 16
    config.model.lstm_layers = 1
    config.training.epochs = args.epochs
    config.training.batch_size = 8
    config.training.beta_warmup_epochs = 1
    rng = np.random.default_rng(config.seed)
    time = np.linspace(0, config.data.seconds, config.data.samples, endpoint=False)
    signals, labels = [], []
    for index in range(32):
        label = index % 2
        frequency = 1.0 if label == 0 else 1.7
        base = np.sin(2 * np.pi * frequency * time)
        leads = np.stack([(1 + lead / 30) * base for lead in range(config.data.leads)])
        signals.append(leads + rng.normal(0, 0.08, leads.shape))
        labels.append(label)
    pipeline, _ = CAREECGPipeline.train(
        np.asarray(signals, dtype=np.float32),
        np.asarray(labels),
        ["synthetic_slow_pattern", "synthetic_fast_pattern"],
        bundled_demo_documents(),
        config,
        validation_fraction=0.125,
    )
    pipeline.save(args.output)
    result = pipeline.run(
        np.asarray(signals[-1], dtype=np.float32),
        "What pattern does this synthetic ECG show?",
        counterfactual_target="synthetic_slow_pattern",
    )
    print(json.dumps(result.to_dict(), indent=2))
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="care-ecg", description="CARE-ECG pipeline")
    subparsers = parser.add_subparsers(dest="command", required=True)

    train = subparsers.add_parser("train", help="train VAE, discretizer, graph, and RAG index")
    train.add_argument("dataset", help="NPZ with signals and labels arrays")
    train.add_argument("output", help="artifact output directory")
    train.add_argument("--config")
    train.add_argument("--source-rate", type=int)
    train.add_argument(
        "--preprocessed",
        action="store_true",
        help="assert that NPZ signals already satisfy the configured preprocessing contract",
    )
    train.add_argument("--label-names", help="comma-separated names for integer labels")
    train.add_argument("--knowledge", help="knowledge documents in JSONL format")
    train.add_argument(
        "--knowledge-graph",
        help="validated graph JSON; retrieval edges are indexed and reviewed priors seed K2",
    )
    train.add_argument("--prior-edges", help="JSON list of [parent, child] clinical-prior edges")
    train.add_argument("--validation-fraction", type=float, default=0.1)
    train.add_argument(
        "--validation-dataset",
        help="optional held-out NPZ (for example PTB-XL fold 9); never used to fit bins or graph",
    )
    train.add_argument("--validation-source-rate", type=int)
    train.set_defaults(func=_train)

    infer = subparsers.add_parser("infer", help="run the full saved pipeline on one ECG")
    infer.add_argument("artifacts")
    infer.add_argument("signal", help="NPY array [12,T] or [T,12]")
    infer.add_argument("query")
    infer.add_argument("--source-rate", type=int)
    infer.add_argument("--patient-id")
    infer.add_argument("--counterfactual-target")
    infer.add_argument("--history", help="NPZ containing prior ECG signals and optional metadata")
    infer.add_argument("--history-source-rate", type=int)
    infer.add_argument("--openai", action="store_true")
    infer.add_argument("--generator", choices=("template", "openai", "huggingface"))
    infer.add_argument("--model", help="model ID for OpenAI or Hugging Face generation")
    infer.add_argument("--output")
    infer.set_defaults(func=_infer)

    predict = subparsers.add_parser(
        "predict", help="run a saved pipeline over an NPZ dataset and write JSONL"
    )
    predict.add_argument("artifacts")
    predict.add_argument("dataset")
    predict.add_argument("output")
    predict.add_argument("--query", default="Which diagnosis is most supported, and why?")
    predict.add_argument("--source-rate", type=int)
    predict.add_argument("--preprocessed", action="store_true")
    predict.add_argument("--counterfactual-target")
    predict.add_argument("--generator", choices=("template", "openai", "huggingface"))
    predict.add_argument("--model")
    predict.add_argument("--limit", type=int)
    predict.add_argument(
        "--batch-size",
        type=int,
        default=64,
        help="neural encoder batch size (default: 64)",
    )
    predict.add_argument(
        "--expert-probabilities",
        metavar="EXPERT_NPZ",
        help=(
            "NPZ with probabilities [N,C], record_ids [N], and label_order [C]; "
            "record order must exactly match the prediction dataset"
        ),
    )
    predict.add_argument(
        "--expert-weight",
        type=float,
        help="validation-selected external-expert weight in [0,1]",
    )
    predict.add_argument(
        "--expert-source",
        help="stable model/checkpoint identifier recorded in every output row",
    )
    predict.set_defaults(func=_predict)

    evaluate = subparsers.add_parser("evaluate", help="evaluate JSONL prediction records")
    evaluate.add_argument("predictions")
    evaluate.add_argument("--scp-aliases", help="JSON mapping SCP codes to text aliases")
    evaluate.add_argument("--expected-field", default="expected")
    evaluate.add_argument("--text-field", default="rationale")
    evaluate.add_argument("--scp-fuzzy-threshold", type=float, default=0.84)
    evaluate.set_defaults(func=_evaluate)

    export_results = subparsers.add_parser(
        "export-results", help="export predictions and paper comparison to an Excel workbook"
    )
    export_results.add_argument("predictions", help="JSONL prediction records")
    export_results.add_argument("output", help="output XLSX path")
    export_results.add_argument("--metadata", help="optional run metadata JSON object")
    export_results.add_argument("--paper-targets", help="optional metric-to-target JSON object")
    export_results.add_argument("--scp-aliases", help="SCP aliases for primary text evaluation")
    export_results.add_argument("--scp-fuzzy-threshold", type=float, default=0.84)
    export_results.add_argument("--tolerance", type=float, default=0.01)
    export_results.set_defaults(func=_export_results)

    prepare = subparsers.add_parser(
        "prepare-ptbxl",
        help="build official fold 1-8/9/10 NPZ splits with five diagnostic superclasses",
    )
    prepare.add_argument("root", help="extracted PTB-XL directory")
    prepare.add_argument("output", help="output directory for train/validation/test NPZ files")
    prepare.add_argument("--sampling-rate", type=int, choices=(100, 500), default=500)
    prepare.add_argument("--seconds", type=int, default=10)
    prepare.add_argument(
        "--normalization",
        choices=("none", "robust", "zscore"),
        default="none",
        help="record-level preprocessing; use none with train-fold global normalization",
    )
    prepare.add_argument("--min-likelihood", type=float, default=0.0)
    prepare.set_defaults(func=_prepare_ptbxl)

    security = subparsers.add_parser(
        "security-run",
        aliases=("security-evaluate",),
        help="run the auditable E0 clean baseline and E1 adaptive waveform attack matrix",
    )
    security.add_argument("artifacts", help="waveform-callable saved CARE artifact directory")
    security.add_argument("dataset", help="raw-mV NPZ with signals, labels, folds, and patient IDs")
    security.add_argument("output", help="new directory for JSON, Excel, logs, and waveforms")
    security.add_argument("--security-config", default="config/security_ptbxl.yaml")
    security.add_argument("--phase", choices=("validation", "heldout"))
    security.add_argument(
        "--validation-lock",
        help="lock created by security-lock; mandatory for held-out evaluation",
    )
    security.add_argument("--source-rate", type=int)
    security.add_argument("--target-ids", help="comma-separated subset of T1,T2,T3,T4")
    security.add_argument("--attacks", help="comma-separated attack implementations")
    security.add_argument("--constraints", help="comma-separated waveform constraints")
    security.add_argument("--budgets", help="comma-separated complete-system query budgets")
    security.add_argument("--seeds", help="comma-separated random seeds")
    security.add_argument("--defense", choices=("v1", "v2"))
    security.add_argument(
        "--max-candidates-per-target",
        type=int,
        help="optional deterministic candidate cap for development or smoke runs",
    )
    security.add_argument("--max-eligible-per-target", type=int)
    security.add_argument(
        "--allow-missing-folds",
        action="store_true",
        help="permit an exploratory run without fold metadata; recorded in the result",
    )
    security.add_argument(
        "--allow-record-level",
        action="store_true",
        help="use record IDs when patient IDs are unavailable; not a primary patient-level run",
    )
    security.add_argument(
        "--excel",
        action=argparse.BooleanOptionalAction,
        default=None,
        help="enable or disable the concise Excel export",
    )
    security.set_defaults(func=_security_evaluate)

    security_lock = subparsers.add_parser(
        "security-lock",
        help="freeze a security config and validation result before held-out evaluation",
    )
    security_lock.add_argument("security_config")
    security_lock.add_argument("validation_results")
    security_lock.add_argument("output")
    security_lock.set_defaults(func=_security_lock)

    security_verify = subparsers.add_parser(
        "security-verify-log",
        help="verify the hash chain and saved waveform hashes in a security event log",
    )
    security_verify.add_argument("log")
    security_verify.add_argument("--waveform-dir")
    security_verify.set_defaults(func=_security_verify_log)

    demo = subparsers.add_parser("demo", help="train and run a small, synthetic offline demo")
    demo.add_argument("--output", default="artifacts/demo")
    demo.add_argument("--epochs", type=int, default=1)
    demo.add_argument("--seed", type=int, default=42)
    demo.set_defaults(func=_demo)
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    return int(args.func(args))


if __name__ == "__main__":
    raise SystemExit(main())
